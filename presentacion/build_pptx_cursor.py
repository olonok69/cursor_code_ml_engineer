#!/usr/bin/env python
"""Genera `presentacion/Cursor_Presentacion.pptx` — el deck del curso de Cursor.

    pip install python-pptx pillow
    python presentacion/build_pptx_cursor.py

36 slides 16:9 en tres partes (1 Cursor · 2 metodología · 3 grafo de tickets).
La slide 34 incrusta `presentacion/kg_graph.png`, que produce `capture_kg_graph.py`.

El fichero está organizado en tres capas:

1. **Sistema visual** — paleta, tipografías y geometría compartida (`CHROME`).
2. **Componentes** — `rect`, `tb`, `card`, `code_block`, `quote`… Cada uno reproduce
   un patrón que se repite por todo el deck; las medidas son las del diseño original.
3. **Contenido** — una función `slide_NN()` por diapositiva. Para editar el deck se
   toca ahí: el texto vive en las llamadas a los componentes, no en los componentes.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "Cursor_Presentacion.pptx"
KG_IMAGE = HERE / "kg_graph.png"

# --------------------------------------------------------------------------- #
# 1. Sistema visual
# --------------------------------------------------------------------------- #

BG = "0F141A"       # fondo de slide
PANEL = "1A212B"    # tarjetas / paneles
BORDER = "2E3948"   # borde de panel
CODE_BG = "0A0E13"  # bloques de código
BAND = "222B38"     # banda de cita y realces
TEXT = "E8EAED"     # texto principal
MUTED = "9AA4B2"    # texto secundario
DIM = "66707E"      # pie de página
ACCENT = "D97757"   # coral — Parte 1 y gates
BLUE = "7AA2C7"     # Parte 2
GREEN = "8CC28C"    # Parte 3 y "esto está bien"

UI = "Segoe UI"
MONO = "Consolas"

SLIDE_W, SLIDE_H = 13.333, 7.5
FOOTER = "Cursor · Curso / Workshop"

THIN = Pt(0.75)  # borde normal
THICK = Pt(1.0)  # borde de énfasis


def In(v):
    """Pulgadas → EMU, aceptando fracciones."""
    return Emu(int(round(v * 914400)))


def R(text, size, color, bold=False, font=UI):
    """Especificación de un run: (texto, cuerpo, color, negrita, tipografía)."""
    return (text, size, color, bold, font)


# --------------------------------------------------------------------------- #
# 2. Componentes
# --------------------------------------------------------------------------- #


def rect(slide, x, y, w, h, fill, border=None, border_w=THIN):
    """Rectángulo sin sombra. `fill=None` => transparente."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, In(x), In(y), In(w), In(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = RGBColor.from_string(border)
        sh.line.width = border_w
    sh.shadow.inherit = False
    return sh


def tb(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
       sa=4, ls=1.0):
    """Caja de texto sin márgenes. `paras` es una lista de listas de runs `R(...)`."""
    box = slide.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sa)
        p.space_before = Pt(0)
        p.line_spacing = ls
        for text, size, color, bold, font in runs:
            run = p.add_run()
            run.text = text
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)
    return box


def line(slide, x, y, w, h, runs, **kw):
    """Atajo para una caja de texto de un solo párrafo."""
    return tb(slide, x, y, w, h, [runs], **kw)


# --- armazón de slide ------------------------------------------------------- #


def new_slide(prs, banded=False):
    """Fondo + cromo lateral. `banded` = franjas superior/inferior (portada y
    separadores de parte); si no, barra vertical izquierda (slides de contenido)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    rect(slide, -0.1, -0.1, 13.53, SLIDE_H + 0.2, BG)
    if banded:
        rect(slide, 0, 0, SLIDE_W, 0.16, ACCENT)
        rect(slide, 0, SLIDE_H - 0.16, SLIDE_W, 0.16, ACCENT)
    else:
        rect(slide, 0, 0, 0.22, SLIDE_H, ACCENT)
    return slide


def page_number(slide, n):
    line(slide, 11.8, 7.02, 1.2, 0.35, [R(str(n), 9.5, DIM, True)], align=PP_ALIGN.RIGHT)


def chrome(prs, eyebrow, title, page):
    """Slide de contenido: barra izquierda, antetítulo, título, regla y pie."""
    slide = new_slide(prs)
    line(slide, 0.7, 0.55, 11.0, 0.4, [R(eyebrow, 12, ACCENT, True)])
    line(slide, 0.68, 0.92, 12.0, 1.1, [R(title, 30, TEXT, True)])
    rect(slide, 0.72, 1.72, 1.1, 0.045, ACCENT)
    line(slide, 0.7, 7.02, 8.0, 0.35, [R(FOOTER, 9.5, DIM)])
    page_number(slide, page)
    return slide


def part_divider(prs, part, title, subtitle, chips, page):
    """Separador de parte: franjas, número de parte, título grande y chips de índice."""
    slide = new_slide(prs, banded=True)
    line(slide, 1.0, 1.3, 11.0, 0.5, [R(part, 16, ACCENT, True)])
    line(slide, 0.95, 1.8, 11.5, 1.2, [R(title, 44, TEXT, True)])
    rect(slide, 1.02, 2.85, 1.4, 0.05, ACCENT)
    line(slide, 1.0, 3.05, 11.3, 0.6, [R(subtitle, 15, MUTED)], ls=1.1)
    for i, text in enumerate(chips):
        x = 1.0 if i % 2 == 0 else 7.0
        y = 3.85 + 0.62 * (i // 2)
        rect(slide, x, y, 5.7, 0.52, PANEL, BORDER)
        line(slide, x + 0.25, y + 0.05, 5.3, 0.42, [R(text, 12.5, TEXT)],
             anchor=MSO_ANCHOR.MIDDLE)
    page_number(slide, page)
    return slide


# --- bloques reutilizables -------------------------------------------------- #


def quote(slide, text, y=6.35):
    """Banda de cierre 🗣️ al pie de la slide."""
    rect(slide, 0.7, y, 11.95, 0.52, BAND, BORDER)
    rect(slide, 0.7, y, 0.09, 0.52, ACCENT)
    line(slide, 0.95, y + 0.03, 11.6, 0.46,
         [R("“ ", 13, ACCENT, True), R(text, 12.5, TEXT), R(" ”", 13, ACCENT, True)],
         anchor=MSO_ANCHOR.MIDDLE)


def panel(slide, x, y, w, h, border=BORDER, border_w=THIN, fill=PANEL):
    return rect(slide, x, y, w, h, fill, border, border_w)


def bullets(slide, x, y, w, h, items, color, size=11.5, ls=1.05):
    """Lista con marcador ▸ del color de la sección."""
    paras = [[R("▸  ", size, color, True), R(t, size, TEXT)] for t in items]
    return tb(slide, x, y, w, h, paras, sa=6, ls=ls)


def card(slide, x, y, w, h, title, items, color, items_h, size=11.5):
    """Panel con encabezado de color + lista de viñetas."""
    panel(slide, x, y, w, h)
    line(slide, x + 0.3, y + 0.22, w - 0.5, 0.4, [R(title, 14, color, True)])
    bullets(slide, x + 0.3, y + 0.72, w - 0.55, items_h, items, color, size)


def num_card(slide, x, y, w, h, badge, title, desc, desc_h):
    """Tarjeta con distintivo cuadrado (número, letra o icono) + título + descripción."""
    panel(slide, x, y, w, h)
    rect(slide, x + 0.22, y + 0.22, 0.5, 0.5, ACCENT)
    line(slide, x + 0.22, y + 0.24, 0.5, 0.46, [R(badge, 15, BG, True)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    line(slide, x + 0.9, y + 0.16, w - 1.05, 0.4, [R(title, 14.5, TEXT, True)])
    line(slide, x + 0.9, y + 0.5, w - 1.05, desc_h, [R(desc, 11, MUTED)], ls=1.02)


def code_block(slide, x, y, w, h, lines):
    """Bloque de código: caja oscura + líneas monoespaciadas.

    `lines` es una lista de cadenas o de tuplas `(texto, color)` — el color se usa
    para los comentarios en gris.
    """
    rect(slide, x, y, w, h, CODE_BG, BORDER)
    paras = []
    for item in lines:
        text, color = item if isinstance(item, tuple) else (item, TEXT)
        paras.append([R(text, 11.5, color, font=MONO)])
    tb(slide, x + 0.28, y + 0.2, w - 0.5, h - 0.35, paras, sa=2, ls=1.12)


def label(slide, x, y, w, text, color=BLUE, size=11):
    """Rótulo pequeño encima de un bloque."""
    return line(slide, x, y, w, 0.3, [R(text, size, color, True)])


def stack_row(slide, y, name, desc, note, x=0.7, w=7.6, h=0.48):
    """Fila apilada (anatomía del contexto): nombre · descripción · nota a la derecha."""
    panel(slide, x, y, w, h)
    line(slide, x + 0.25, y + 0.04, 2.6, 0.4, [R(name, 11.5, BLUE, True)],
         anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 3.6, y + 0.04, 4.0, 0.4, [R(desc, 10.5, TEXT)], anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 7.35, y + 0.04, 0.85, 0.4, [R(note, 9.0, MUTED)],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def flag_row(slide, y, verdict, headline, detail):
    """Fila ancha con pestaña de color a la izquierda (verde = bien, coral = ojo)."""
    color = GREEN if verdict else ACCENT
    panel(slide, 0.7, y, 11.95, 0.58)
    rect(slide, 0.7, y, 0.08, 0.58, color)
    line(slide, 0.95, y + 0.05, 4.6, 0.48, [R(headline, 11.5, color, True)],
         anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 5.7, y + 0.05, 6.8, 0.48, [R(detail, 11.5, TEXT)],
         anchor=MSO_ANCHOR.MIDDLE)


def step_row(slide, y, label_text, detail):
    """Fila del ejemplo real: etapa + qué pasó."""
    panel(slide, 0.7, y, 11.95, 0.48)
    rect(slide, 0.7, y, 0.08, 0.48, ACCENT)
    line(slide, 0.95, y + 0.04, 2.4, 0.4, [R(label_text, 11.5, BLUE, True)],
         anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 3.5, y + 0.04, 9.0, 0.4, [R(detail, 11, TEXT)], anchor=MSO_ANCHOR.MIDDLE)


def tool_row(slide, y, phase, tool, note, highlight=False):
    """Fila de prevalencia de tools: fase · herramienta · coste."""
    if highlight:
        panel(slide, 0.7, y, 11.95, 0.5, border=ACCENT, border_w=THICK, fill=BAND)
    else:
        panel(slide, 0.7, y, 11.95, 0.5)
    line(slide, 0.95, y + 0.05, 2.9, 0.4,
         [R(phase, 11.5, ACCENT if highlight else BLUE, True)], anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 3.95, y + 0.05, 7.0, 0.4, [R(tool, 11.5, TEXT)], anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 11.1, y + 0.05, 1.4, 0.4, [R(note, 9.5, MUTED)],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def compare_row(slide, y, field, left, right):
    """Fila de la tabla comparativa subagent vs Background/Cloud Agent."""
    panel(slide, 0.7, y, 11.95, 0.52)
    line(slide, 0.95, y + 0.05, 3.3, 0.42, [R(field, 11.5, TEXT, True)],
         anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 4.6, y + 0.05, 3.7, 0.42, [R(left, 11, MUTED)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 8.6, y + 0.05, 3.7, 0.42, [R(right, 11, MUTED)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def flow_step(slide, x, y, number, text, gate=False):
    """Etapa del flujo de 11 pasos. Las etapas coral son gates."""
    panel(slide, x, y, 5.85, 0.56,
          border=ACCENT if gate else BORDER, border_w=THICK if gate else THIN)
    rect(slide, x + 0.14, y + 0.11, 0.34, 0.34, ACCENT if gate else BAND)
    line(slide, x + 0.14, y + 0.11, 0.34, 0.34,
         [R(number, 11, BG if gate else TEXT, True)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    line(slide, x + 0.62, y + 0.06, 5.1, 0.46, [R(text, 11.5, TEXT)],
         anchor=MSO_ANCHOR.MIDDLE)


def summary_card(slide, x, y, number, title, desc, color):
    """Tarjeta compacta del cierre."""
    panel(slide, x, y, 5.85, 0.85)
    rect(slide, x + 0.2, y + 0.18, 0.44, 0.44, color)
    line(slide, x + 0.2, y + 0.19, 0.44, 0.42, [R(number, 13, BG, True)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    line(slide, x + 0.8, y + 0.1, 5.0, 0.35, [R(title, 12.5, TEXT, True)])
    line(slide, x + 0.8, y + 0.43, 5.0, 0.38, [R(desc, 10, MUTED)])


# --------------------------------------------------------------------------- #
# 3. Contenido — una función por slide
# --------------------------------------------------------------------------- #


def slide_01(prs):
    """Portada."""
    slide = new_slide(prs, banded=True)
    line(slide, 1.0, 1.4, 11.0, 0.5,
         [R("CURSO / WORKSHOP · DEVELOPER TOOLING", 14, ACCENT, True)])
    tb(slide, 0.95, 2.0, 11.5, 2.0, [
        [R("Cursor", 60, TEXT, True)],
        [R("la misma metodología, ", 26, MUTED), R("otra superficie", 26, TEXT, True)],
    ], sa=10)
    rect(slide, 1.02, 4.2, 1.4, 0.05, ACCENT)
    tb(slide, 1.0, 4.45, 11.5, 1.7, [
        [R("PARTE 1 — ", 15, ACCENT, True),
         R("Cursor: instalación · AGENTS.md & rules · contexto & caching · MCP · "
           "skills & marketplace · subagents · automatización", 13.5, MUTED)],
        [R("PARTE 2 — ", 15, BLUE, True),
         R("La metodología (agnóstica): flujo con gates · CodeGraph / Serena / GSD · "
           "transferencia · ops (machine-sync + S3)", 13.5, MUTED)],
        [R("PARTE 3 — ", 15, GREEN, True),
         R("El grafo de conocimiento de tickets, construido con graphify: pipeline, "
           "skills kg / kg-refresh y el grafo real", 13.5, MUTED)],
    ], sa=8)
    line(slide, 1.0, 6.7, 11.0, 0.4,
         [R("Volumen 2 del curso: mismo material que la versión Claude Code, verificado "
            "contra cursor.com/docs y reescrito para Cursor (agosto 2026).", 11, DIM)])


def slide_02(prs):
    """Agenda."""
    slide = chrome(prs, "AGENDA", "Un curso, tres partes", 2)

    def column(x, y0, heading, color, items):
        line(slide, x, y0, 5.85, 0.4, [R(heading, 14, color, True)])
        for i, text in enumerate(items):
            y = y0 + 0.45 + 0.56 * i
            panel(slide, x, y, 5.85, 0.48)
            line(slide, x + 0.25, y + 0.04, 5.5, 0.4, [R(text, 12, TEXT)],
                 anchor=MSO_ANCHOR.MIDDLE)

    column(0.7, 1.95, "PARTE 1 · Cursor", ACCENT, [
        "01  Instalación y uso básico",
        "02  Memoria, instrucciones y sesiones",
        "03  Contexto: context window y prompt caching",
        "04  MCP — conectar tus herramientas",
        "05  Skills y Marketplace",
        "06  Subagents",
        "07  Automatización",
    ])
    column(6.8, 1.95, "PARTE 2 · La metodología (agnóstica)", BLUE, [
        "08  El flujo de 11 etapas + ejemplo real",
        "09  Las herramientas del método",
        "10  Transferir: de Claude Code a Cursor",
        "11  Ops: machine-sync + registro S3",
    ])
    column(6.8, 4.75, "PARTE 3 · El grafo de tickets (graphify)", GREEN, [
        "12  Diseño · pipeline · skills kg / kg-refresh · el grafo real",
        "13  Cierre",
    ])


def slide_03(prs):
    part_divider(
        prs, "PARTE 1", "Cursor",
        "La herramienta: del editor en tu escritorio a agentes en la nube.",
        ["01 · Instalación y uso básico", "02 · Memoria y sesiones",
         "03 · Contexto y prompt caching", "04 · MCP",
         "05 · Skills y Marketplace", "06 · Subagents",
         "07 · Automatización"],
        3)


def slide_04(prs):
    slide = chrome(prs, "PARTE 1 · 01 · INSTALACIÓN Y USO BÁSICO",
                   "Tres formas de entrar: editor, CLI, nube", 4)
    cards = [
        ("1", "Editor Cursor",
         "Descarga desde cursor.com — macOS / Windows / Linux."),
        ("2", "Cursor CLI (agent)",
         "curl … | bash (Unix) · irm … | iex (Windows) — headless-capable."),
        ("3", "Background / Cloud Agents",
         "cursor.com/agents — corre tareas en la nube sin editor abierto."),
        ("4", "Multi-repo / workspace",
         "Abre varios repos; distingue propiedad de contexto de solo-lectura."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.05 + 1.28 * (i // 2)
        num_card(slide, x, y, 5.835, 1.0, badge, title, desc, 0.45)

    label(slide, 0.7, 4.51, 5.9, "Arrancar en cualquier proyecto")
    code_block(slide, 0.7, 4.85, 5.9, 1.3, [
        "cd tu-proyecto",
        "agent               # CLI interactivo, login la 1ª vez",
    ])
    card(slide, 6.9, 4.51, 5.75, 1.55, "El mismo agente en todas partes", [
        "Editor (chat / Agent mode) · CLI (agent)",
        "Web (cursor.com/agents, tareas async)",
        "Bugbot revisa PRs automáticamente en GitHub",
    ], BLUE, 0.65)
    quote(slide, "Instalar es trivial; el valor está en cómo lo usas.")


def slide_05(prs):
    slide = chrome(prs, "PARTE 1 · 01 · INSTALACIÓN Y USO BÁSICO",
                   "Dos modos: interactivo y headless", 5)
    card(slide, 0.7, 2.05, 5.85, 2.5, "Interactivo — pensar contigo", [
        "Editor (chat/Agent mode) o CLI `agent`",
        "Plan mode: propone un plan y tú lo apruebas antes de tocar nada",
        "Ideal para explorar, diseñar, depurar",
    ], ACCENT, 1.6)
    card(slide, 6.8, 2.05, 5.85, 2.5, "Headless (agent -p) — scripts y CI", [
        "Print mode: prompt in (arg), resultado por stdout",
        "Combina con --force / --output-format text|json",
        "La base de la automatización (sección 07)",
    ], BLUE, 1.6)
    label(slide, 0.7, 4.71, 11.95, "El mismo Cursor, dos formas de invocarlo")
    code_block(slide, 0.7, 5.05, 11.95, 1.1, [
        'agent -p "resume los cambios de esta rama"',
        'agent -p --output-format text "revisa por seguridad los ficheros tocados vs main"',
    ])


def slide_06(prs):
    slide = chrome(prs, "PARTE 1 · 02 · MEMORIA E INSTRUCCIONES",
                   "AGENTS.md: el patrón de dos niveles", 6)
    card(slide, 0.7, 2.05, 5.85, 3.0, "Nivel 1 — siempre cargado (pequeño)", [
        "Orientación mínima: qué es el proyecto, mapa de repos, comandos",
        "Solo PUNTEROS de una línea a todo lo demás",
        "Es la onboarding de 30s del agente",
        "Regla write-once: el core apunta, no copia",
    ], ACCENT, 2.1)
    card(slide, 6.8, 2.05, 5.85, 3.0, "Nivel 2 — bajo demanda (el detalle)", [
        "data/changes/STATUS.md — estado vivo por ticket",
        "PLAYBOOK.md — lecciones de debugging",
        "SHARP_EDGES.md — invariantes 'no tocar'",
        "<TICKET>/<TICKET>.md — ledgers autocontenidos",
    ], BLUE, 2.1)
    quote(slide, "El mismo recorte que en Claude Code: punteros, no copias — "
                 "AGENTS.md se mantiene deliberadamente pequeño.")


def slide_07(prs):
    slide = chrome(prs, "PARTE 1 · 02 · INSTRUCCIONES, PERMISOS Y MEMORIES",
                   "AGENTS.md, rules, permisos y memoria", 7)
    label(slide, 0.7, 1.81, 6.0, "Rules (.mdc) + AGENTS.md (se combinan, no compiten)")
    code_block(slide, 0.7, 2.15, 6.0, 2.35, [
        ".cursor/rules/00-core.mdc     # alwaysApply: true",
        ".cursor/rules/01-tools.mdc    # @fichero = incluye contenido",
        "AGENTS.md                     # raíz (+ anidado: gana el más",
        "                               # específico), sin frontmatter",
        "~/.cursor/permissions.json    # mcpAllowlist server:tool",
    ])
    card(slide, 7.0, 1.82, 5.65, 2.55, "Permisos = allowlist, no solo UI", [
        "permissions.json: mcpAllowlist / terminalAllowlist (glob)",
        "4 modos de rule: Always / Apply Intelligently /",
        "  Apply to Specific Files / Apply Manually",
        "El humano es dueño de push/PR/deploy",
    ], GREEN, 1.65)
    card(slide, 0.7, 4.6, 11.95, 1.5,
         "Memories — un sistema DISTINTO de AGENTS.md/rules", [
             "Cursor genera Memories automáticamente a partir de tus chats "
             "(no son ficheros que tú escribes).",
             "No es el mismo mecanismo que la auto-memory de Claude Code: trátalas "
             "como complemento, no como sustituto del patrón de dos niveles.",
         ], BLUE, 0.6)


def slide_08(prs):
    slide = chrome(prs, "PARTE 1 · 02 · SESIONES",
                   "Dónde vive Cursor: editor, CLI y nube", 8)
    cards = [
        ("▣", "Editor",
         "Chat + Agent mode + Plan mode; diffs inline, revisión en el propio IDE."),
        ("›_", "Cursor CLI",
         "agent en terminal — igual que el editor, sin UI; ideal para SSH/servers."),
        ("☁", "Background / Cloud Agents",
         "cursor.com/agents corre tareas async; sigue el progreso desde la web."),
        ("✓", "Bugbot",
         "Revisa cada PR en GitHub/GitLab/Bitbucket automáticamente, sin script propio."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.15 + 1.38 * (i // 2)
        num_card(slide, x, y, 5.835, 1.1, badge, title, desc, 0.55)
    quote(slide, "Mismo agente, mismos AGENTS.md/rules/MCP — cambia solo dónde se ejecuta.")


def slide_09(prs):
    slide = chrome(prs, "PARTE 1 · 03 · CONTEXTO",
                   "Context window: el recurso que gobierna todo", 9)
    rows = [
        ("System / agent prompt", "Oculto — siempre primero", "fijo"),
        ("Rules alwaysApply + AGENTS.md",
         "Lo controlas tú — por eso el patrón de dos niveles", "tú decides"),
        ("Memories", "Si las hay — distinto sistema, revisa qué se coló", "variable"),
        ("Tools MCP", "Schemas de las tools habilitadas", "por server"),
        ("Conversación + resultados",
         "Ficheros leídos, output de comandos… crece cada turno", "crece"),
    ]
    for i, (name, desc, note) in enumerate(rows):
        stack_row(slide, 2.0 + 0.58 * i, name, desc, note)

    label(slide, 8.6, 2.01, 4.05, "Los mandos (CLI + UI del editor)")
    code_block(slide, 8.6, 2.35, 4.05, 2.55, [
        "agent> /summarize  # liberar contexto",
        "agent> /rewind     # mensaje previo",
        ("# nueva chat / nueva invocación de agent", MUTED),
        ("# + anillo de contexto en el editor", MUTED),
    ])
    line(slide, 0.7, 5.05, 11.95, 0.9, [
        R("El rendimiento degrada ANTES de llenar la ventana: ", 12, TEXT, True),
        R("contexto con ruido = peores decisiones. Cursor resume automáticamente al "
          "acercarse al límite (además del /summarize manual) — no asumas paridad "
          "exacta con /compact de Claude Code.", 12, MUTED),
    ], ls=1.1)
    quote(slide, "El anillo de contexto del editor te dice qué bloque engorda — "
                 "mide antes de recortar.")


def slide_10(prs):
    slide = chrome(prs, "PARTE 1 · 03 · CONTEXTO",
                   "Higiene de contexto: máxima señal por token", 10)
    cards = [
        ("1", "AGENTS.md / rules mínimos",
         "Si puedes borrarlo sin que el agente se equivoque, bórralo. "
         "Dos niveles + punteros."),
        ("2", "@fichero solo si hace falta",
         "Las rules .mdc soportan @fichero para incluir contenido — úsalo con cuidado, "
         "es carga eager."),
        ("3", "MCP con moderación",
         "Cada server suma su bloque de tools. Desactiva los que el proyecto no use."),
        ("4", "Subagents para investigar",
         "La exploración sucia va a un subagent (sección 06); vuelve solo el resumen."),
        ("5", "Plan → Agent",
         "Separa exploración (Plan mode) de implementación; el ruido no se queda a vivir."),
        ("6", "Lecturas con puntería",
         "“lee config/auth.ts” > “entiende el auth”. Adelanto: CodeGraph-primero "
         "(Parte 2)."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.0 + 1.46 * (i // 2)
        num_card(slide, x, y, 5.835, 1.18, badge, title, desc, 0.63)
    quote(slide, "La prevalencia de tools de la Parte 2 es, en el fondo, "
                 "política de contexto.")


def slide_11(prs):
    slide = chrome(prs, "PARTE 1 · 03 · CONTEXTO",
                   "Prompt caching: no pagar lo mismo dos veces", 11)
    line(slide, 0.7, 1.9, 11.95, 0.55, [
        R("Esto es de la API de Anthropic, no de Cursor — pero si automatizas con la "
          "Cursor SDK contra Claude, aplica igual. Orden estricto ", 12.5, MUTED),
        R("Tools → System → Messages", 12.5, ACCENT, True, MONO),
        R(": un cambio invalida su nivel y los siguientes.", 12.5, MUTED),
    ])
    for i, (ttl, write, read) in enumerate([
        ("TTL 5 min (defecto)", "escribir 1.25×", "leer 0.1×"),
        ("TTL 1 hora", "escribir 2×", "leer 0.1×"),
    ]):
        y = 2.6 + 0.6 * i
        panel(slide, 0.7, y, 5.85, 0.5)
        line(slide, 0.95, y + 0.05, 2.5, 0.4, [R(ttl, 11.5, BLUE, True)],
             anchor=MSO_ANCHOR.MIDDLE)
        line(slide, 3.5, y + 0.05, 1.7, 0.4, [R(write, 11.5, TEXT)],
             anchor=MSO_ANCHOR.MIDDLE)
        line(slide, 5.2, y + 0.05, 1.2, 0.4, [R(read, 11.5, GREEN, True)],
             anchor=MSO_ANCHOR.MIDDLE)

    tb(slide, 0.7, 3.95, 5.85, 1.6, [
        [R("Una sesión de 50 turnos relee el prefijo 50 veces a 0.1× — ", 12, TEXT),
         R("ese es el descuento que hace viables las sesiones largas.", 12, GREEN, True)],
        [R("Mínimo cacheable ~1k tokens · máx. 4 breakpoints · diagnóstico en "
           "usage.cache_read_input_tokens.", 10.5, MUTED)],
    ], sa=8, ls=1.1)

    label(slide, 6.8, 2.26, 5.85, "Lo estable primero; el breakpoint al final de lo estable")
    code_block(slide, 6.8, 2.6, 5.85, 2.9, [
        'system=[{ "type": "text",',
        '  "text": STABLE_INSTRUCTIONS,   # estable',
        '  "cache_control":',
        '    {"type": "ephemeral"} }]   # breakpoint',
        'messages=[{"role": "user",',
        '  "content": query }]  # variable: FUERA',
        ("# error clásico: timestamp ANTES del", MUTED),
        ("# breakpoint -> 0 hits y nadie sabe por qué", MUTED),
    ])
    quote(slide, "Demo ejecutable: ejemplos/prompt-caching/cache_demo.py — habla con la "
                 "API Anthropic, no con el producto Cursor.")


def slide_12(prs):
    slide = chrome(prs, "PARTE 1 · 03 · CONTEXTO",
                   "Caching: qué controlas en Cursor y qué no", 12)
    rows = [
        (True, "AGENTS.md + rules pequeños y estables",
         "Prefijo corto que no cambia entre turnos → mejor comportamiento"),
        (False, "Editar rules/AGENTS.md a mitad de sesión",
         "Cambia el contexto base → paga impuesto de nuevo"),
        (False, "Muchos servers MCP activos",
         "Bloque de tools grande y cambiante → más contexto fijo"),
        (True, "Nueva sesión / chat limpio entre tareas no relacionadas",
         "Evita arrastrar transcript infinito"),
        (False, "Asumir que Cursor expone cache_control como la API",
         "No — es un producto distinto; no copies env vars de Claude Code"),
    ]
    for i, (verdict, headline, detail) in enumerate(rows):
        flag_row(slide, 2.05 + 0.68 * i, verdict, headline, detail)
    line(slide, 0.7, 5.6, 11.95, 0.55, [
        R("Cursor no expone cache_control como producto: aplica lo que el proveedor del "
          "modelo decida. Lo que sí controlas: tamaño/estabilidad de AGENTS.md + rules "
          "+ tools MCP.", 11.5, MUTED)])
    quote(slide, "Contexto lean y estable rinde mejor en cualquier producto — Cursor "
                 "incluido — aunque no veas el descuento en pantalla.")


def slide_13(prs):
    slide = chrome(prs, "PARTE 1 · 04 · MCP",
                   "Model Context Protocol: conecta tu mundo", 13)
    line(slide, 0.7, 1.9, 11.95, 0.6, [
        R("Mismo estándar abierto que en Claude Code — CodeGraph, Serena, Playwright "
          "hablan MCP en ambos. Config de proyecto en ", 12.5, MUTED),
        R(".cursor/mcp.json", 12.5, ACCENT, True, MONO),
        R(".", 12.5, MUTED),
    ])
    cards = [
        ("P", "Scope project", ".cursor/mcp.json — versionado, compartido con el equipo."),
        ("U", "Scope user",
         "~/.cursor/mcp.json — editable directamente, o vía Settings → MCP."),
        ("✓", "Permisos", "permissions.json: mcpAllowlist / terminalAllowlist."),
        ("★", "Servers del día a día", "serena · context7 · playwright · codegraph (+ supabase opcional)."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.785
        y = 2.55 + 1.07 * (i // 2)
        num_card(slide, x, y, 5.865, 0.85, badge, title, desc, 0.3)

    label(slide, 0.7, 4.62, 11.95,
          "Añadir un server (secretos por env var; recarga obligatoria)")
    code_block(slide, 0.7, 4.96, 11.95, 1.1, [
        '"serena": { "command": "uvx", "args": ["--from", '
        '"git+https://github.com/oraios/serena", "serena", "start-mcp-server"] }',
        ("# tras editar .cursor/mcp.json: recarga/reinicia Cursor — no hay hot-reload",
         MUTED),
    ])


def slide_14(prs):
    slide = chrome(prs, "PARTE 1 · 05 · SKILLS Y MARKETPLACE",
                   "Skills: compatible con Claude Code", 14)
    cards = [
        ("1", "Tools",
         "Lo que el agente puede hacer: Read/Edit/Shell/Grep + Task + mcp__*. "
         "Gobernadas por permissions.json (mcpAllowlist)."),
        ("2", "Skills",
         ".cursor/skills/<n>/SKILL.md con frontmatter name+description. "
         "Auto-selección por description."),
        ("3", "Interop con Claude Code",
         "Cursor también lee .claude/skills/ — el MISMO SKILL.md sirve en los dos "
         "productos."),
        ("4", "Marketplace",
         "cursor.com/marketplace: paquetes instalables de skills+subagents+MCP+hooks"
         "+rules."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.05 + 1.43 * (i // 2)
        num_card(slide, x, y, 5.835, 1.15, badge, title, desc, 0.6)

    line(slide, 0.7, 4.75, 11.95, 0.9, [
        R("Lo que YA NO es una brecha frente a Claude Code: ", 12.5, GREEN, True),
        R("cuando se escribió la guía original de adaptación no existían ni Skills ni "
          "Marketplace en Cursor. Hoy ambos son reales — revisa docs.cursor.com antes "
          "de asumir que algo “no tiene equivalente”.", 12.5, MUTED),
    ], ls=1.05)
    quote(slide, "Skill = capacidad que Cursor decide usar por su description. "
                 "Marketplace = el reparto versionado.")


def slide_15(prs):
    slide = chrome(prs, "PARTE 1 · 06 · SUBAGENTS",
                   "Subagents: el ruido muere fuera de tu sesión", 15)
    card(slide, 0.7, 2.0, 5.85, 2.2, "Feature nativa de Cursor", [
        "Delegación tipo Task: contexto aislado por subagent",
        "Se invoca por lenguaje natural (o /nombre)",
        "Ejecución en paralelo para trabajo independiente",
        "Tipos base documentados de forma laxa — no asumas nombres exactos",
    ], BLUE, 1.3)
    label(slide, 6.8, 2.01, 5.85, 'Custom = .cursor/agents/<nombre>.md')
    code_block(slide, 6.8, 2.35, 5.85, 2.2, [
        ("# .cursor/agents/refactor-scout.md", MUTED),
        "---",
        "name: refactor-scout",
        "description: Scout blast radius before rename",
        "readonly: true",
        "---",
        "Usa CodeGraph y LUEGO Serena…",
    ])
    card(slide, 0.7, 4.75, 5.85, 1.5, "Lo que hay que saber", [
        "Contexto AISLADO: solo el resumen vuelve a tu sesión",
        "NO hereda tu conversación: dale contexto en el prompt de lanzamiento",
    ], ACCENT, 0.6)
    quote(slide, "Ejemplo listo: ejemplos/subagents/.cursor/agents/refactor-scout.md "
                 "(también prompts/ si prefieres pegar al lanzar).")


def slide_16(prs):
    slide = chrome(prs, "PARTE 1 · 06 · SUBAGENTS",
                   "Paralelismo real: varios Background/Cloud Agents", 16)
    card(slide, 0.7, 2.0, 5.85, 2.6, "Lo que NO existe en Cursor", [
        "Agent Teams de Claude Code: lead + teammates + inbox compartido",
        "Sin mensajería directa entre agentes",
        "No lo intentes portar 1:1 — no hay producto equivalente",
    ], ACCENT, 1.7)
    label(slide, 6.8, 1.81, 5.85, "El sustituto pragmático")
    code_block(slide, 6.8, 2.15, 5.85, 1.9, [
        "// cursor.com/agents — varias tareas en paralelo",
        "// cada Background/Cloud Agent = sesión completa,",
        "// en su propia rama, sin coordinarse entre sí",
    ])
    card(slide, 6.8, 4.25, 5.85, 1.9, "Cómo se usa en la práctica", [
        "Particiona el trabajo por rama/fichero antes de lanzar",
        "Un humano (o el agente principal) integra resultados",
        "Coste: cada Cloud Agent es una sesión completa",
    ], BLUE, 1.0)
    line(slide, 0.7, 4.75, 5.85, 1.3, [
        R("Se pide en lenguaje natural: ", 11.5, MUTED),
        R("lanza tres Background Agents, uno por módulo, cada uno en su propia rama; "
          "revisa y haz merge tú.", 11.5, TEXT),
    ], ls=1.1)
    quote(slide, "Sin inbox compartido: particiona los ficheros/ramas — cada agente es "
                 "dueño de los suyos.")


def slide_17(prs):
    slide = chrome(prs, "PARTE 1 · 06 · SUBAGENTS",
                   "¿Subagent o Background/Cloud Agent? La decisión", 17)
    line(slide, 4.6, 2.0, 3.7, 0.4, [R("SUBAGENT", 13, BLUE, True)],
         align=PP_ALIGN.CENTER)
    line(slide, 8.6, 2.0, 3.7, 0.4, [R("BACKGROUND / CLOUD AGENT", 13, ACCENT, True)],
         align=PP_ALIGN.CENTER)
    rows = [
        ("Contexto", "Aislado; devuelve un resumen", "Sesión completa, async, en su rama"),
        ("Comunicación", "Solo resultado → sesión principal",
         "Ninguna entre agentes (sin inbox)"),
        ("Coste", "Bajo (lo caro muere fuera)", "Alto (N sesiones completas)"),
        ("Úsalo para", "Side-quests: investigar, verificar",
         "Trabajo largo/async, o en paralelo real"),
        ("Config", ".cursor/agents/*.md, prompt o skill",
         "cursor.com/agents (UI / handoff &)"),
    ]
    for i, (field, left, right) in enumerate(rows):
        compare_row(slide, 2.45 + 0.62 * i, field, left, right)
    line(slide, 0.7, 5.68, 11.95, 0.5, [
        R("Puente a la Parte 2: ", 11.5, GREEN, True),
        R("GSD (Claude Code) empaqueta roles como subagentes plugin; en Cursor esos "
          "roles viven como .cursor/agents/ + skills — este proyecto usa data/changes/, "
          "no GSD (ver Parte 2).", 11.5, MUTED),
    ])
    quote(slide, "Subagent para que el ruido muera fuera; Background/Cloud Agent para "
                 "trabajo largo o async. El coste no es el mismo.")


def slide_18(prs):
    slide = chrome(prs, "PARTE 1 · 07 · AUTOMATIZACIÓN",
                   "Hooks: el control determinista", 18)
    card(slide, 0.7, 2.0, 5.85, 3.1, "El contrato de un hook", [
        ".cursor/hooks.json (proyecto) o ~/.cursor/hooks.json (usuario)",
        "Payload del evento por STDIN, respuesta JSON por STDOUT",
        "Eventos: beforeShellExecution, beforeMCPExecution, beforeReadFile,",
        "  afterFileEdit, preToolUse/postToolUse, beforeSubmitPrompt, stop…",
        "permission: prefer deny para gates (ask a menudo se ignora)",
        "exit 2 = deny también · failClosed si el hook crashea",
    ], ACCENT, 2.2)
    label(slide, 6.8, 2.01, 5.85, "Patrón de bloqueo (JS)")
    code_block(slide, 6.8, 2.35, 5.85, 2.4, [
        "// beforeReadFile — bloquear .env",
        'const p = JSON.parse(require("fs")',
        '  .readFileSync(0, "utf8"));',
        'if ((p.path || "").includes(".env")) {',
        "  console.log(JSON.stringify(",
        '    {permission:"deny",',
        '     user_message:"Bloqueado: .env"}));',
        "  process.exit(0);",
        "}",
    ])
    quote(slide, "Con AGENTS.md le pides que se porte bien; con un hook lo garantizas — "
                 "igual que en Claude Code, distinto contrato.")


def slide_19(prs):
    slide = chrome(prs, "PARTE 1 · 07 · AUTOMATIZACIÓN",
                   "De hooks a agentes en la nube", 19)
    cards = [
        ("a", "Hooks",
         "Seguridad (.env), formato, type-check bloqueante, veto de push/deploy."),
        ("b", "Headless / CLI",
         "agent -p en scripts y CI (print mode; --force / --output-format)."),
        ("c", "CI/CD",
         "Bugbot (nativo) para revisión de PR, o Cursor SDK en tu propio GitHub Action."),
        ("d", "Scheduling",
         "Automations: cron + triggers (Slack, Linear, PR merged, PagerDuty)."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.785
        y = 2.05 + 1.17 * (i // 2)
        num_card(slide, x, y, 5.865, 0.95, badge, title, desc, 0.4)

    label(slide, 0.7, 4.34, 11.95,
          "Cursor SDK (@cursor/sdk) — Agent.prompt (one-shot) · Agent.create+send (stream/multi-turn)")
    code_block(slide, 0.7, 4.68, 11.95, 1.3, [
        'import { Agent } from "@cursor/sdk";',
        "const result = await Agent.prompt(prompt, {",
        "  apiKey: process.env.CURSOR_API_KEY!, local: { cwd } });",
        "// multi-turn: Agent.create(...) → agent.send(...) → run.stream()",
    ])


def slide_20(prs):
    part_divider(
        prs, "PARTE 2", "La metodología",
        "Agnóstica de la herramienta: se demuestra con Cursor, pero la disciplina viaja.",
        ["08 · El flujo de 11 etapas + ejemplo real",
         "09 · Las herramientas del método",
         "10 · Transferir: de Claude Code a Cursor",
         "11 · Ops: machine-sync + registro compartido (S3)"],
        20)


def slide_21(prs):
    slide = chrome(prs, "PARTE 2 · 08 · LA METODOLOGÍA",
                   "El agente es un colaborador disciplinado", 21)
    panel(slide, 0.7, 2.0, 11.95, 1.15, border=ACCENT, border_w=THICK, fill=BAND)
    line(slide, 1.0, 2.18, 11.4, 0.85,
         [R("“La autonomía se gana por-decisión, no se concede en bloque.”", 17, TEXT,
            True)],
         anchor=MSO_ANCHOR.MIDDLE)
    card(slide, 0.7, 3.4, 5.85, 2.05, "El agente posee", [
        "Investigación y diagnóstico",
        "Planes e implementación",
        "Tests y documentación",
    ], BLUE, 1.15)
    card(slide, 6.8, 3.4, 5.85, 2.05, "El humano posee", [
        "Decisiones go/no-go y el scope",
        "Toda acción externa: push, PR, deploy",
        "La aprobación del plan",
    ], GREEN, 1.15)
    line(slide, 0.72, 5.62, 11.9, 0.5, [
        R("Agnóstica:  ", 11.5, ACCENT, True),
        R("no es un flujo de Cursor — es una forma de trabajar con CUALQUIER agente. "
          "La sección 10 lo demuestra transfiriéndola desde Claude Code (y también "
          "llega a GitHub Copilot).", 11, MUTED),
    ], ls=1.03)
    quote(slide, "El modelo no gana confianza gratis: la gana decisión a decisión, "
                 "con evidencia.")


def slide_22(prs):
    slide = chrome(prs, "PARTE 2 · 08 · LA METODOLOGÍA", "El flujo real de 11 etapas", 22)
    steps = [
        ("1", "Orientar: skill kg + history-first + status-first", False),
        ("2", "Triaje inbound: ¿síntoma en el contrato?", True),
        ("3", "Regresión vs. pre-existente (probarlo)", True),
        ("4", "Investigar: oráculo determinista barato", False),
        ("5", "Plan mode → acuerdo humano explícito", True),
        ("6", "Implementar: TDD RED → GREEN, mínimo", False),
        ("7", "Verificar: gate outbound ×5 (instrumento → lista → imagen → mirar)", True),
        ("8", "Documentar — cada cosa una vez", False),
        ("9", "Sanitizar — líneas añadidas (skill sanitise-diff)", False),
        ("10", "Handoff: el humano hace push / PR", False),
        ("11", "Revisión Bugbot/bot + persistir lecciones", True),
    ]
    for i, (number, text, gate) in enumerate(steps):
        x = 0.7 if i < 6 else 6.85
        y = 2.05 + 0.66 * (i if i < 6 else i - 6)
        flow_step(slide, x, y, number, text, gate)
    line(slide, 0.7, 6.02, 11.95, 0.3, [
        R("■ ", 11, ACCENT, True),
        R("Los recuadros coral son GATES (puntos de decisión). Un gate rojo = STOP: "
          "no escribir código.", 11, MUTED),
    ])
    quote(slide, "El agente orquesta y es donde vive la inferencia; lo caro se concentra "
                 "en plan/código/verify, no en buscar.")


def slide_23(prs):
    slide = chrome(prs, "PARTE 2 · 08 · LA METODOLOGÍA",
                   "Un ejemplo real, de principio a fin", 23)
    panel(slide, 0.7, 1.95, 11.95, 0.62, border=ACCENT, fill=BAND)
    line(slide, 0.95, 2.0, 11.5, 0.52, [
        R("Bug QA: ", 12.5, ACCENT, True),
        R("“un campo sale vacío en la UI, pero está escrito en el PDF.”", 12.5, TEXT),
    ], anchor=MSO_ANCHOR.MIDDLE)
    rows = [
        ("Orientar",
         "Skill kg saca la zona de peligro (un SHARP_EDGE que restringe el fix); "
         "git/gh dan la base."),
        ("Contrato",
         "El campo está vacío en el JSON del endpoint (Playwright MCP) → el Lambda es "
         "responsable."),
        ("Provenance",
         "Re-extraer en la base previa: ya salía vacío → pre-existente, no regresión."),
        ("Investigar",
         "Serena + CodeGraph (MCP) localizan el detector; un _diag_pdf.py determinista "
         "da la causa: SIN LLM."),
        ("Fix + verificar",
         "Test RED → fix estructural (no cliente) → regresión no-op + contrato local "
         "(wrapper) + dentro de la imagen desplegada."),
        ("Handoff",
         "Skill sanitise-diff → el humano hace push/PR. Bugbot detecta un caso más → "
         "test + PLAYBOOK."),
    ]
    for i, (label_text, detail) in enumerate(rows):
        step_row(slide, 2.78 + 0.56 * i, label_text, detail)
    quote(slide, "El modelo no gana confianza gratis: la gana decisión a decisión, "
                 "con evidencia.")


def slide_24(prs):
    slide = chrome(prs, "PARTE 2 · 09 · HERRAMIENTAS DEL MÉTODO",
                   "Qué herramienta usa el agente, y cuándo", 24)
    line(slide, 0.7, 1.9, 11.95, 0.5, [
        R("Regla de .cursor/rules/01-tool-prevalence.mdc: ", 12, MUTED),
        R("codegraph_explore PRIMERO (survey en 1 llamada) · Serena "
          "find_referencing_symbols para el chequeo preciso antes de renombrar. ",
          12, TEXT, True),
        R("Orden: barato → caro, determinista → probabilístico.", 12, MUTED),
    ])
    rows = [
        ("Orientar",
         "Skill kg (grafo de tickets → zona de peligro) · STATUS.md · git · gh",
         "sin inferencia", False),
        ("Navegar (survey)",
         "CodeGraph codegraph_explore (MCP) — fuente + rutas + blast radius + cobertura",
         "1 llamada", False),
        ("Refactor-check",
         "Serena find_referencing_symbols (MCP) — desambigua por clase "
         "(antes de renombrar)", "preciso", False),
        ("Diagnosticar", "Oráculo determinista: parser / validador / _diag_*.py",
         "sin inferencia", False),
        ("Entorno (logs, config)",
         "AWS CLI — CloudWatch · lambda get-function · SQS/DLQ", "read-only", False),
        ("Contrato de salida", "Playwright MCP sobre el endpoint que ve el consumidor",
         "reproducir", False),
        ("Solo al final",
         "La tirada del agente — para VERIFICAR el fix, no para diagnosticar",
         "metered", True),
    ]
    for i, (phase, tool, note, highlight) in enumerate(rows):
        tool_row(slide, 2.5 + 0.6 * i, phase, tool, note, highlight)


def slide_25(prs):
    slide = chrome(prs, "PARTE 2 · 09 · HERRAMIENTAS DEL MÉTODO",
                   "Cada herramienta, en una frase — y su fase", 25)
    cards = [
        ("CG", "CodeGraph → investigar",
         "Grafo del CÓDIGO (tree-sitter→SQLite, local, vía MCP). Survey en 1 llamada: "
         "fuente + callers + blast radius + cobertura."),
        ("Se", "Serena → pre-refactor",
         "Navegación semántica vía LSP (MCP). find_referencing_symbols desambigua por "
         "clase: el chequeo preciso antes de renombrar."),
        ("G", "GSD → plan/ejecutar/verificar",
         "El método productizado — solo Claude Code hoy. En Cursor: Plan mode + skill "
         "methodology-plan + data/changes/."),
        ("kg", "Skill kg → orientar",
         "Grafo de la MEMORIA del proyecto (con graphify): tickets + sharp edges. "
         "History-first sin LLM (Parte 3)."),
        ("Pw", "Playwright (MCP) → triaje y outbound",
         "Reproduce el síntoma donde lo ve el consumidor: el endpoint real, no una "
         "función interna."),
        ("Or", "Oráculos → diagnosticar",
         "Parsers/validadores/_diag_*.py propios: respuesta barata y reproducible antes "
         "de gastar la tirada del agente."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.0 + 1.46 * (i // 2)
        num_card(slide, x, y, 5.835, 1.18, badge, title, desc, 0.63)
    quote(slide, "La inversión clásica — el modelo diagnostica — es lo que este orden "
                 "evita: el modelo verifica; los oráculos diagnostican.")


def slide_26(prs):
    slide = chrome(prs, "PARTE 2 · 09 · HERRAMIENTAS DEL MÉTODO",
                   "GSD: solo existe en Claude Code (hoy)", 26)
    card(slide, 0.7, 2.0, 5.85, 3.3, "Qué es GSD", [
        "Plugin de Claude Code para gestión de proyectos por fases",
        "Ciclo: discuss → plan (gate) → execute → verify",
        "Subagentes gsd-planner, gsd-plan-checker, gsd-executor,",
        "  gsd-code-reviewer, gsd-verifier",
        "Estado versionado en .planning/",
        "No hay port oficial a Cursor",
    ], ACCENT, 2.4)
    card(slide, 6.8, 2.0, 5.85, 3.3, "Equivalente en Cursor", [
        "Plan mode para el gate discuss → plan",
        "Skill methodology-plan rellena la plantilla antes de implementar",
        "Roles gsd-* → Task/subagents + prompts con ese rol",
        "Estado en data/changes/ (igual que Claude Code)",
        "GSD brilla en greenfield multi-fase; aquí no se usa a diario",
    ], BLUE, 2.4)
    quote(slide, "Mismo método, sin el plugin: en Cursor, Plan mode + skills replican "
                 "los gates de GSD.")


def slide_27(prs):
    slide = chrome(prs, "PARTE 2 · 09 · HERRAMIENTAS DEL MÉTODO",
                   "CodeGraph: inteligencia de código local (vía MCP)", 27)
    card(slide, 0.7, 2.0, 5.85, 2.4, "Índice tree-sitter → SQLite (.codegraph/)", [
        "Símbolos: funciones, clases, rutas, componentes",
        "Aristas: llamadas, imports, herencia, referencias",
        "Determinista (del AST), sin API keys",
        "Devuelve: fuente + rutas + blast radius + cobertura",
    ], BLUE, 1.5)
    label(slide, 6.8, 1.86, 5.85, "CLI + registro en Cursor (no 'claude mcp add')")
    code_block(slide, 6.8, 2.2, 5.85, 1.85, [
        "codegraph init      # crea .codegraph/",
        "codegraph sync      # incremental tras editar",
        'codegraph explore "<símbolo|pregunta>"',
        ("# .cursor/mcp.json:", MUTED),
        '"codegraph": {"command":"codegraph",',
        '  "args":["serve","--path","<repo>","--mcp"]}',
    ])
    panel(slide, 6.8, 4.25, 5.85, 1.1)
    line(slide, 7.05, 4.38, 5.4, 0.95, [
        R("Primero para navegar:  ", 11, ACCENT, True),
        R("fuente + callers + blast radius + cobertura en 1 consulta (trátala como YA "
          "leída). Serena find_referencing_symbols = chequeo preciso (desambigua por "
          "clase).", 11, TEXT),
    ], ls=1.03)
    tb(slide, 0.7, 4.6, 5.85, 1.4, [
        [R("58% menos tool calls · 22% más rápido", 13, GREEN, True)],
        [R("(según sus benchmarks: casi elimina las lecturas de fichero)", 10.5, MUTED)],
    ])
    quote(slide, "Una consulta en vez de grep→abrir→seguir-import→repetir. "
                 "Menos contexto, más señal.")


def slide_28(prs):
    slide = chrome(prs, "PARTE 2 · 09 · HERRAMIENTAS DEL MÉTODO",
                   "Serena: navegación semántica, chequeo preciso", 28)
    card(slide, 0.7, 2.0, 5.85, 2.6, "Navegación semántica vía LSP (MCP)", [
        "Opera sobre SÍMBOLOS, no texto: precisión de IDE",
        "find_symbol (body=true) — un método de un fichero de 5k líneas",
        "get_symbols_overview — esqueleto de un fichero",
        "find_referencing_symbols — quién referencia, POR CLASE",
        "search_for_pattern · activate_project (multi-repo)",
    ], BLUE, 1.7)
    label(slide, 6.8, 2.01, 5.85, "Instalar (stdio) — no uses 'claude mcp add'")
    code_block(slide, 6.8, 2.35, 5.85, 1.35, [
        "// .cursor/mcp.json",
        '"serena": {"command":"uvx","args":["--from",',
        '  "git+https://github.com/oraios/serena",',
        '  "serena","start-mcp-server"]}',
    ])
    panel(slide, 6.8, 4.0, 5.85, 1.55, border=ACCENT)
    line(slide, 7.05, 4.15, 5.4, 1.3, [
        R("Por qué es OBLIGATORIO pre-rename:  ", 11, ACCENT, True),
        R("el impact plano de CodeGraph mezcla métodos homónimos (Invoice.process vs "
          "Refund.process); Serena los desambigua por clase. Complementarios, "
          "no rivales.", 11, TEXT),
    ], ls=1.05)
    line(slide, 0.7, 4.85, 5.85, 0.9, [
        R("La plantilla refactor-scout ", 11.5, GREEN, True),
        R("(ejemplos/subagents/prompts/) empaqueta el orden CodeGraph → Serena → grep "
          "como procedimiento.", 11.5, MUTED),
    ], ls=1.05)
    quote(slide, "CodeGraph responde '¿qué se rompe?'; Serena responde "
                 "'¿exactamente quién llama a ESTE process()?'")


def slide_29(prs):
    slide = chrome(prs, "PARTE 2 · 10 · TRANSFERENCIA",
                   "La prueba real: de Claude Code a Cursor", 29)
    card(slide, 0.7, 2.0, 5.85, 2.75, "Viaja SIN cambios (las 5 reglas)", [
        "Plan → acuerdo → implementar",
        "Verificar en el contrato del CONSUMIDOR",
        "Resolver la clase general, no un input",
        "Rastro durable: porqué, qué, cómo se verificó",
        "El humano posee lo externo (merge, deploy)",
    ], GREEN, 1.85)
    card(slide, 6.8, 2.0, 5.85, 2.75, "Lo que SÍ cambia: la superficie", [
        "CLAUDE.md (+ jerarquía) → AGENTS.md + .cursor/rules/*.mdc",
        "Skills ~/.claude/skills/ → .cursor/skills/ (mismo SKILL.md)",
        "Hooks + settings.local.json → .cursor/hooks.json",
        "Plan mode → Plan mode (misma disciplina)",
        "Subagents/Agent Teams → .cursor/agents/*.md (sin Teams)",
        "claude -p (headless) → agent -p (Cursor CLI print mode)",
    ], BLUE, 1.85)
    card(slide, 0.7, 4.95, 11.95, 1.25,
         "El starter-kit: CURSOR_ADAPTATION.md (docs/ai-agents-code-methodology/)", [
             "El mapeo completo + cursor/ (rules, skills, MCP, hooks listos para copiar) "
             "+ bootstrap-cursor-repo.ps1",
             "El mismo pack trae COPILOT_ADAPTATION.md: no es un caso especial — la "
             "disciplina llega a un tercer agente",
         ], ACCENT, 0.35)
    quote(slide, "Este mismo repo es la prueba: metodología nacida en Claude Code, "
                 "corriendo en Cursor con CURSOR_ADAPTATION.md.")


def slide_30(prs):
    slide = chrome(prs, "PARTE 2 · 11 · OPS",
                   "De transportar (tarball) a compartir (S3)", 30)
    card(slide, 0.7, 2.0, 5.85, 2.55, "A · Bring-up: tarball + USB", [
        "Outbound = copia completa; inbound = solo delta de data/",
        "USB: mount manual en WSL + verificar byte a byte",
        "En destino: bootstrap + target-setup.sh rehacen el tooling",
        "Sigue siendo el camino para levantar una máquina desde cero",
    ], BLUE, 1.65)
    card(slide, 6.8, 2.0, 5.85, 2.55, "B · Día a día: registro sobre S3", [
        "Alcance estrecho: changes/**/*.md + grafo (nada de cliente)",
        "Escribe por sync; lee por mount de SOLO LECTURA",
        "Dry-run por defecto; --delete opt-in (no borrar al compañero)",
        "UN solo publisher. \"Derivado\" es del fichero, no de la carpeta",
    ], GREEN, 1.65)
    card(slide, 0.7, 4.7, 11.95, 1.5,
         "Lo específico de agentes: IDENTITY.md machine-local", [
             "Cada máquina declara MACHINE_NAME + MACHINE_ROLE (publisher | contributor)",
             "AGENTS.md apunta a IDENTITY.md — la sesión lee su rol ANTES de actuar "
             "(si no, un contributor republica el grafo)",
         ], ACCENT, 0.55)
    quote(slide, "Cuando el rastro durable pasa de una máquina a un equipo, el agente "
                 "tiene que saber en qué máquina está antes de actuar.")


def slide_31(prs):
    part_divider(
        prs, "PARTE 3", "El grafo de tickets",
        "Un caso completo construido con graphify — la memoria del proyecto, navegable "
        "desde Cursor.",
        ["12a · El problema y el diseño (spike)",
         "12b · Pipeline y comandos (skill kg, kg-refresh)",
         "12c · El grafo real, visualizado",
         "12d · Uso y enganche en la metodología"],
        31)


def slide_32(prs):
    slide = chrome(prs, "PARTE 3 · 12A · GRAFO DE TICKETS",
                   "El problema, y por qué graphify", 32)
    card(slide, 0.7, 2.0, 5.85, 2.5, "El problema", [
        "~540 ficheros de writeups, sharp edges, runbooks, memoria",
        "Un bug 'nuevo' casi siempre tiene contexto previo que restringe el fix",
        "Encontrarlo a mano = recordar que existe + grep",
        "El grafo lo hace EXPLÍCITO y consultable en 1 llamada",
    ], ACCENT, 1.6)
    card(slide, 6.8, 2.0, 5.85, 2.5, "El corpus: manifest, no glob", [
        "manifest.txt DIFFEABLE: ~116 ficheros, ~196k palabras",
        "Writeups sst-* (+ fallback determinista) · hubs · runbooks · memoria",
        "Exclusiones duras: binarios, handovers repetidos, copias stale (payload/)",
        "Densidad sin conocimiento nuevo = ruido",
    ], BLUE, 1.6)
    panel(slide, 0.7, 4.75, 11.95, 1.35, border=GREEN, fill=BAND)
    line(slide, 1.0, 4.92, 11.4, 1.05, [
        R("La tecnología es graphify — CodeGraph es solo la analogía.  ", 12.5, GREEN,
          True),
        R("Mismo rol (grafo consultable antes de tocar nada), otro dominio (tickets, no "
          "código), otra herramienta. El pipeline de build se generó originalmente en "
          "Claude Code; hoy la consulta y el refresco viven como skills kg / kg-refresh "
          "— el mismo SKILL.md funciona en Cursor y en Claude Code.", 11.5, TEXT),
    ], ls=1.12)
    quote(slide, "Material completo (diseño, scripts, tests, salida real): "
                 "docs/knowledge-graph/.")


def slide_33(prs):
    slide = chrome(prs, "PARTE 3 · 12B · GRAFO DE TICKETS",
                   "Pipeline y skills — kg y kg-refresh", 33)
    label(slide, 0.7, 1.91, 6.35, "Build (skill kg-refresh)")
    code_block(slide, 0.7, 2.25, 6.35, 2.15, [
        "kg_refresh.sh prepare   # manifest -> stage",
        "                        # -> scratch FUERA del repo",
        ("# extracción semántica con el agente:", MUTED),
        "                        # nodos+aristas+clustering",
        "kg_refresh.sh finalize  # output/ + leak-check",
    ])
    card(slide, 7.35, 1.92, 5.3, 2.5, "Las piezas", [
        "Skills kg y kg-refresh (.cursor/skills/ o .claude/skills/ — mismo fichero)",
        "kg_query.sh — envuelve graphify explain/path + find",
        "kg_refresh.sh — prepare · finalize · bootstrap · snapshot/restore-memory",
        "build_manifest.py + stage_corpus.py — el corpus",
        "test_kg_*.py — los bookends, TESTEADOS",
    ], BLUE, 1.6)
    panel(slide, 0.7, 4.62, 11.95, 1.5)
    tb(slide, 1.0, 4.78, 11.4, 1.2, [
        [R("El gotcha que lo sostiene:  ", 12, ACCENT, True),
         R("graphify respeta .gitignore y todo data/ lo está → correr el detector in "
           "situ encuentra 0 ficheros. Por eso el corpus se monta en un scratch fuera "
           "del repo y los artefactos se copian de vuelta.", 11.5, TEXT)],
        [R("Por eso kg-refresh es un skill, no solo un script:  ", 12, GREEN, True),
         R("el paso semántico (extracción, subagentes en paralelo) necesita un agente; "
           "los bookends son deterministas.", 11.5, TEXT)],
    ], sa=8, ls=1.12)
    quote(slide, "Staging con provenance: sst-5468__sst-5468.md, hub__STATUS.md, "
                 "memory__x.md — cada nodo traza a su fuente.")


def slide_34(prs):
    slide = chrome(prs, "PARTE 3 · 12C · GRAFO DE TICKETS",
                   "El grafo real: 507 nodos, 35 comunidades", 34)
    rect(slide, 0.68, 1.93, 8.84, 4.99, None, BORDER, THICK)
    if not KG_IMAGE.exists():
        raise SystemExit(
            f"Falta {KG_IMAGE.name}: genéralo con "
            "`python presentacion/capture_kg_graph.py` (requiere playwright)."
        )
    slide.shapes.add_picture(str(KG_IMAGE), In(0.7), In(1.95), In(8.8), In(4.95))

    card(slide, 9.7, 1.95, 2.95, 3.4, "La salida real", [
        "507 nodos · 672 aristas",
        "35 comunidades",
        "92% EXTRACTED (fiable)",
        "7% INFERRED (conf. 0.7)",
        "116 ficheros, ~196k palabras",
    ], GREEN, 2.5)
    tb(slide, 9.7, 5.5, 2.95, 0.8, [
        [R("graph.html interactivo:", 11, BLUE, True)],
        [R("búsqueda de nodos, filtro por comunidad (vis-network).", 10, MUTED)],
    ], ls=1.05)
    quote(slide, "Las comunidades mapean a zonas de peligro reales; los god-nodes son "
                 "la lista de onboarding gratis.")


def slide_35(prs):
    slide = chrome(prs, "PARTE 3 · 12D · GRAFO DE TICKETS",
                   "Cómo se usa — y dónde se engancha", 35)
    label(slide, 0.7, 1.91, 6.35, "Consulta: CERO LLM (kg_query.sh lee graph.json)")
    code_block(slide, 0.7, 2.25, 6.35, 2.0, [
        "kg_query.sh explain <ticket|tema>  # vecinos",
        "kg_query.sh path <A> <B>           # camino",
        "kg_query.sh find <substr>          # nombre exacto",
        ("# skill kg-refresh                 # reconstruir (barato)", MUTED),
    ])
    card(slide, 7.35, 1.92, 5.3, 2.5, "Dónde se engancha", [
        "Etapa 1 (Orientar) del flujo de 11 etapas",
        "Regla history-first de .cursor/rules/00-methodology-core.mdc:",
        "skill kg <ticket|tema> ANTES de grep",
        "Apunta a QUÉ leer; no lo sustituye",
    ], ACCENT, 1.6)
    line(slide, 0.7, 4.45, 6.35, 0.9, [
        R("Ejemplo real: kg get_letter_end ", 12, GREEN, True),
        R("→ la zona de peligro completa al instante: los 5-6 tickets que comparten ese "
          "código.", 11.5, MUTED),
    ], ls=1.1)
    card(slide, 0.7, 5.2, 11.95, 1.0, "Honestidad y ciclo de vida", [
        "Recall en zonas densas · EXTRACTED = fiable, INFERRED = pista · interno "
        "(data/) · derivado se reconstruye, pero lo escrito a mano viaja",
    ], BLUE, 0.1)
    quote(slide, "Un paso semántico en el build, cero LLM en la consulta. El grafo es "
                 "el mapa; el agente, el guía.")


def slide_36(prs):
    slide = chrome(prs, "CIERRE", "De asistente a sistema, en cualquier editor", 36)
    line(slide, 0.7, 1.95, 5.85, 0.4, [R("PARTE 1 · la herramienta", 13, ACCENT, True)])
    left = [
        ("1", "Instalar + memoria",
         "Editor, CLI (agent) o nube; AGENTS.md + rules de dos niveles; Memories aparte."),
        ("2", "Contexto & caching",
         "El presupuesto y el descuento: lean y estable gana en ambos, aunque no lo veas "
         "en pantalla."),
        ("3", "MCP + skills + Marketplace",
         "Conecta tu mundo; empaqueta y distribuye flujos — y comparte SKILL.md con "
         "Claude Code."),
        ("4", "Subagents + automatización",
         "Escala el trabajo; hooks que garantizan calidad; Bugbot y Automations en la "
         "nube."),
    ]
    for i, (number, title, desc) in enumerate(left):
        summary_card(slide, 0.7, 2.35 + 0.95 * i, number, title, desc, ACCENT)

    line(slide, 6.8, 1.95, 5.85, 0.4,
         [R("PARTE 2 + 3 · el método y el grafo", 13, BLUE, True)])
    right = [
        ("5", "El flujo de 11 etapas",
         "Gates deterministas; el humano posee las decisiones."),
        ("6", "Las tools del método",
         "CodeGraph · Serena · oráculos: barato→caro; GSD = el método productizado "
         "(Claude Code)."),
        ("7", "Transferible y hasta en ops",
         "CURSOR_ADAPTATION.md; machine-sync + registro S3 con identidad por máquina."),
        ("8", "El grafo de tickets (graphify)",
         "507 nodos · 35 comunidades: history-first sin LLM, skill kg en cualquier "
         "agente."),
    ]
    for i, (number, title, desc) in enumerate(right):
        summary_card(slide, 6.8, 2.35 + 0.95 * i, number, title, desc, BLUE)

    line(slide, 0.7, 6.35, 12.0, 0.4, [
        R("Referencias:  cursor.com/docs  ·  docs.cursor.com/hooks  ·  "
          "cursor.com/marketplace  ·  github.com/colbymchenry/codegraph  ·  "
          "github.com/oraios/serena", 11, DIM)])


SLIDES = [
    slide_01, slide_02, slide_03, slide_04, slide_05, slide_06, slide_07, slide_08,
    slide_09, slide_10, slide_11, slide_12, slide_13, slide_14, slide_15, slide_16,
    slide_17, slide_18, slide_19, slide_20, slide_21, slide_22, slide_23, slide_24,
    slide_25, slide_26, slide_27, slide_28, slide_29, slide_30, slide_31, slide_32,
    slide_33, slide_34, slide_35, slide_36,
]


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    for make in SLIDES:
        make(prs)
    prs.save(str(OUT))
    return len(SLIDES)


if __name__ == "__main__":
    n = build()
    print(f"OK  {OUT}  ({n} slides)")
