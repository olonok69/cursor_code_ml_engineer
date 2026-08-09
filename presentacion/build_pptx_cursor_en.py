#!/usr/bin/env python
"""Build `presentacion/Cursor_Presentacion_EN.pptx` — English Cursor course deck.

    pip install python-pptx pillow
    python presentacion/build_pptx_cursor_en.py

36 slides 16:9 in three parts (1 Cursor · 2 methodology · 3 ticket knowledge graph).
Slide 34 embeds `presentacion/kg_graph.png`, produced by `capture_kg_graph.py`.

The file is organized in three layers:

1. **Visual system** — shared palette, type, and geometry (`CHROME`).
2. **Components** — `rect`, `tb`, `card`, `code_block`, `quote`… Each reproduces
   a pattern used across the deck; measurements match the original design.
3. **Content** — one `slide_NN()` function per slide. To edit the deck, change
   those: text lives in component calls, not in the components.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "Cursor_Presentacion_EN.pptx"
KG_IMAGE = HERE / "kg_graph.png"

# --------------------------------------------------------------------------- #
# 1. Visual system
# --------------------------------------------------------------------------- #

BG = "0F141A"       # slide background
PANEL = "1A212B"    # cards / panels
BORDER = "2E3948"   # panel border
CODE_BG = "0A0E13"  # code blocks
BAND = "222B38"     # quote band and highlights
TEXT = "E8EAED"     # primary text
MUTED = "9AA4B2"    # secondary text
DIM = "66707E"      # footer
ACCENT = "D97757"   # coral — Part 1 and gates
BLUE = "7AA2C7"     # Part 2
GREEN = "8CC28C"    # Part 3 and "this is good"

UI = "Segoe UI"
MONO = "Consolas"

SLIDE_W, SLIDE_H = 13.333, 7.5
FOOTER = "Cursor · Course / Workshop"

THIN = Pt(0.75)  # normal border
THICK = Pt(1.0)  # emphasis border


def In(v):
    """Inches → EMU, accepting fractions."""
    return Emu(int(round(v * 914400)))


def R(text, size, color, bold=False, font=UI):
    """Run spec: (text, size, color, bold, font)."""
    return (text, size, color, bold, font)


# --------------------------------------------------------------------------- #
# 2. Components
# --------------------------------------------------------------------------- #


def rect(slide, x, y, w, h, fill, border=None, border_w=THIN):
    """Rectangle with no shadow. `fill=None` => transparent."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, In(x), In(y), In(w), In(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = RGBColor.from_string(border)
        sh.line.width = border_w
    sh.shadow.inherit = False
    return sh


def tb(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
       sa=4, ls=1.0):
    """Text box with no margins. `paras` is a list of lists of `R(...)` runs."""
    box = slide.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sa)
        p.space_before = Pt(0)
        p.line_spacing = ls
        for text, size, color, bold, font in runs:
            run = p.add_run()
            run.text = text
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)
    return box


def line(slide, x, y, w, h, runs, **kw):
    """Shortcut for a single-paragraph text box."""
    return tb(slide, x, y, w, h, [runs], **kw)


# --- slide chrome ----------------------------------------------------------- #


def new_slide(prs, banded=False):
    """Background + side chrome. `banded` = top/bottom stripes (title and
    part dividers); otherwise left vertical bar (content slides)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    rect(slide, -0.1, -0.1, 13.53, SLIDE_H + 0.2, BG)
    if banded:
        rect(slide, 0, 0, SLIDE_W, 0.16, ACCENT)
        rect(slide, 0, SLIDE_H - 0.16, SLIDE_W, 0.16, ACCENT)
    else:
        rect(slide, 0, 0, 0.22, SLIDE_H, ACCENT)
    return slide


def page_number(slide, n):
    line(slide, 11.8, 7.02, 1.2, 0.35, [R(str(n), 9.5, DIM, True)], align=PP_ALIGN.RIGHT)


def chrome(prs, eyebrow, title, page):
    """Content slide: left bar, eyebrow, title, rule, and footer."""
    slide = new_slide(prs)
    line(slide, 0.7, 0.55, 11.0, 0.4, [R(eyebrow, 12, ACCENT, True)])
    line(slide, 0.68, 0.92, 12.0, 1.1, [R(title, 30, TEXT, True)])
    rect(slide, 0.72, 1.72, 1.1, 0.045, ACCENT)
    line(slide, 0.7, 7.02, 8.0, 0.35, [R(FOOTER, 9.5, DIM)])
    page_number(slide, page)
    return slide


def part_divider(prs, part, title, subtitle, chips, page):
    """Part divider: stripes, part number, large title, and index chips."""
    slide = new_slide(prs, banded=True)
    line(slide, 1.0, 1.3, 11.0, 0.5, [R(part, 16, ACCENT, True)])
    line(slide, 0.95, 1.8, 11.5, 1.2, [R(title, 44, TEXT, True)])
    rect(slide, 1.02, 2.85, 1.4, 0.05, ACCENT)
    line(slide, 1.0, 3.05, 11.3, 0.6, [R(subtitle, 15, MUTED)], ls=1.1)
    for i, text in enumerate(chips):
        x = 1.0 if i % 2 == 0 else 7.0
        y = 3.85 + 0.62 * (i // 2)
        rect(slide, x, y, 5.7, 0.52, PANEL, BORDER)
        line(slide, x + 0.25, y + 0.05, 5.3, 0.42, [R(text, 12.5, TEXT)],
             anchor=MSO_ANCHOR.MIDDLE)
    page_number(slide, page)
    return slide


# --- reusable blocks -------------------------------------------------------- #


def quote(slide, text, y=6.35):
    """Closing quote band at the foot of the slide."""
    rect(slide, 0.7, y, 11.95, 0.52, BAND, BORDER)
    rect(slide, 0.7, y, 0.09, 0.52, ACCENT)
    line(slide, 0.95, y + 0.03, 11.6, 0.46,
         [R("“ ", 13, ACCENT, True), R(text, 12.5, TEXT), R(" ”", 13, ACCENT, True)],
         anchor=MSO_ANCHOR.MIDDLE)


def panel(slide, x, y, w, h, border=BORDER, border_w=THIN, fill=PANEL):
    return rect(slide, x, y, w, h, fill, border, border_w)


def bullets(slide, x, y, w, h, items, color, size=11.5, ls=1.05):
    """List with ▸ marker in the section color."""
    paras = [[R("▸  ", size, color, True), R(t, size, TEXT)] for t in items]
    return tb(slide, x, y, w, h, paras, sa=6, ls=ls)


def card(slide, x, y, w, h, title, items, color, items_h, size=11.5):
    """Panel with colored header + bullet list."""
    panel(slide, x, y, w, h)
    line(slide, x + 0.3, y + 0.22, w - 0.5, 0.4, [R(title, 14, color, True)])
    bullets(slide, x + 0.3, y + 0.72, w - 0.55, items_h, items, color, size)


def num_card(slide, x, y, w, h, badge, title, desc, desc_h):
    """Card with square badge (number, letter, or icon) + title + description."""
    panel(slide, x, y, w, h)
    rect(slide, x + 0.22, y + 0.22, 0.5, 0.5, ACCENT)
    line(slide, x + 0.22, y + 0.24, 0.5, 0.46, [R(badge, 15, BG, True)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    line(slide, x + 0.9, y + 0.16, w - 1.05, 0.4, [R(title, 14.5, TEXT, True)])
    line(slide, x + 0.9, y + 0.5, w - 1.05, desc_h, [R(desc, 11, MUTED)], ls=1.02)


def code_block(slide, x, y, w, h, lines):
    """Code block: dark box + monospace lines.

    `lines` is a list of strings or `(text, color)` tuples — color is used
    for gray comments.
    """
    rect(slide, x, y, w, h, CODE_BG, BORDER)
    paras = []
    for item in lines:
        text, color = item if isinstance(item, tuple) else (item, TEXT)
        paras.append([R(text, 11.5, color, font=MONO)])
    tb(slide, x + 0.28, y + 0.2, w - 0.5, h - 0.35, paras, sa=2, ls=1.12)


def label(slide, x, y, w, text, color=BLUE, size=11):
    """Small label above a block."""
    return line(slide, x, y, w, 0.3, [R(text, size, color, True)])


def stack_row(slide, y, name, desc, note, x=0.7, w=7.6, h=0.48):
    """Stacked row (context anatomy): name · description · note on the right."""
    panel(slide, x, y, w, h)
    line(slide, x + 0.25, y + 0.04, 2.6, 0.4, [R(name, 11.5, BLUE, True)],
         anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 3.6, y + 0.04, 4.0, 0.4, [R(desc, 10.5, TEXT)], anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 7.35, y + 0.04, 0.85, 0.4, [R(note, 9.0, MUTED)],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def flag_row(slide, y, verdict, headline, detail):
    """Wide row with colored left tab (green = good, coral = watch out)."""
    color = GREEN if verdict else ACCENT
    panel(slide, 0.7, y, 11.95, 0.58)
    rect(slide, 0.7, y, 0.08, 0.58, color)
    line(slide, 0.95, y + 0.05, 4.6, 0.48, [R(headline, 11.5, color, True)],
         anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 5.7, y + 0.05, 6.8, 0.48, [R(detail, 11.5, TEXT)],
         anchor=MSO_ANCHOR.MIDDLE)


def step_row(slide, y, label_text, detail):
    """Real-example row: stage + what happened."""
    panel(slide, 0.7, y, 11.95, 0.48)
    rect(slide, 0.7, y, 0.08, 0.48, ACCENT)
    line(slide, 0.95, y + 0.04, 2.4, 0.4, [R(label_text, 11.5, BLUE, True)],
         anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 3.5, y + 0.04, 9.0, 0.4, [R(detail, 11, TEXT)], anchor=MSO_ANCHOR.MIDDLE)


def tool_row(slide, y, phase, tool, note, highlight=False):
    """Tool-prevalence row: phase · tool · cost."""
    if highlight:
        panel(slide, 0.7, y, 11.95, 0.5, border=ACCENT, border_w=THICK, fill=BAND)
    else:
        panel(slide, 0.7, y, 11.95, 0.5)
    line(slide, 0.95, y + 0.05, 2.9, 0.4,
         [R(phase, 11.5, ACCENT if highlight else BLUE, True)], anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 3.95, y + 0.05, 7.0, 0.4, [R(tool, 11.5, TEXT)], anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 11.1, y + 0.05, 1.4, 0.4, [R(note, 9.5, MUTED)],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def compare_row(slide, y, field, left, right):
    """Comparative table row: subagent vs Background/Cloud Agent."""
    panel(slide, 0.7, y, 11.95, 0.52)
    line(slide, 0.95, y + 0.05, 3.3, 0.42, [R(field, 11.5, TEXT, True)],
         anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 4.6, y + 0.05, 3.7, 0.42, [R(left, 11, MUTED)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    line(slide, 8.6, y + 0.05, 3.7, 0.42, [R(right, 11, MUTED)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def flow_step(slide, x, y, number, text, gate=False):
    """Step in the 11-stage flow. Coral stages are gates."""
    panel(slide, x, y, 5.85, 0.56,
          border=ACCENT if gate else BORDER, border_w=THICK if gate else THIN)
    rect(slide, x + 0.14, y + 0.11, 0.34, 0.34, ACCENT if gate else BAND)
    line(slide, x + 0.14, y + 0.11, 0.34, 0.34,
         [R(number, 11, BG if gate else TEXT, True)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    line(slide, x + 0.62, y + 0.06, 5.1, 0.46, [R(text, 11.5, TEXT)],
         anchor=MSO_ANCHOR.MIDDLE)


def summary_card(slide, x, y, number, title, desc, color):
    """Compact closing card."""
    panel(slide, x, y, 5.85, 0.85)
    rect(slide, x + 0.2, y + 0.18, 0.44, 0.44, color)
    line(slide, x + 0.2, y + 0.19, 0.44, 0.42, [R(number, 13, BG, True)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    line(slide, x + 0.8, y + 0.1, 5.0, 0.35, [R(title, 12.5, TEXT, True)])
    line(slide, x + 0.8, y + 0.43, 5.0, 0.38, [R(desc, 10, MUTED)])


# --------------------------------------------------------------------------- #
# 3. Content — one function per slide
# --------------------------------------------------------------------------- #


def slide_01(prs):
    """Title slide."""
    slide = new_slide(prs, banded=True)
    line(slide, 1.0, 1.4, 11.0, 0.5,
         [R("COURSE / WORKSHOP · DEVELOPER TOOLING", 14, ACCENT, True)])
    tb(slide, 0.95, 2.0, 11.5, 2.0, [
        [R("Cursor", 60, TEXT, True)],
        [R("the same methodology, ", 26, MUTED), R("a different surface", 26, TEXT, True)],
    ], sa=10)
    rect(slide, 1.02, 4.2, 1.4, 0.05, ACCENT)
    tb(slide, 1.0, 4.45, 11.5, 1.7, [
        [R("PART 1 — ", 15, ACCENT, True),
         R("Cursor: install · AGENTS.md & rules · context & caching · MCP · "
           "skills & marketplace · subagents · automation", 13.5, MUTED)],
        [R("PART 2 — ", 15, BLUE, True),
         R("The methodology (tool-agnostic): gated flow · CodeGraph / Serena / GSD · "
           "transfer · ops (machine-sync + S3)", 13.5, MUTED)],
        [R("PART 3 — ", 15, GREEN, True),
         R("The ticket knowledge graph, built with graphify: pipeline, "
           "skills kg / kg-refresh, and the real graph", 13.5, MUTED)],
    ], sa=8)
    line(slide, 1.0, 6.7, 11.0, 0.4,
         [R("Course volume 2: same material as the Claude Code edition, verified "
            "against cursor.com/docs and rewritten for Cursor (August 2026).", 11, DIM)])


def slide_02(prs):
    """Agenda."""
    slide = chrome(prs, "AGENDA", "One course, three parts", 2)

    def column(x, y0, heading, color, items):
        line(slide, x, y0, 5.85, 0.4, [R(heading, 14, color, True)])
        for i, text in enumerate(items):
            y = y0 + 0.45 + 0.56 * i
            panel(slide, x, y, 5.85, 0.48)
            line(slide, x + 0.25, y + 0.04, 5.5, 0.4, [R(text, 12, TEXT)],
                 anchor=MSO_ANCHOR.MIDDLE)

    column(0.7, 1.95, "PART 1 · Cursor", ACCENT, [
        "01  Install and basic usage",
        "02  Memory, instructions, and sessions",
        "03  Context: context window and prompt caching",
        "04  MCP — connect your tools",
        "05  Skills and Marketplace",
        "06  Subagents",
        "07  Automation",
    ])
    column(6.8, 1.95, "PART 2 · The methodology (tool-agnostic)", BLUE, [
        "08  The 11-stage flow + real example",
        "09  The method's tools",
        "10  Transfer: from Claude Code to Cursor",
        "11  Ops: machine-sync + S3 record",
    ])
    column(6.8, 4.75, "PART 3 · The ticket graph (graphify)", GREEN, [
        "12  Design · pipeline · skills kg / kg-refresh · the real graph",
        "13  Close",
    ])


def slide_03(prs):
    part_divider(
        prs, "PART 1", "Cursor",
        "The tool: from the editor on your desk to agents in the cloud.",
        ["01 · Install and basic usage", "02 · Memory and sessions",
         "03 · Context and prompt caching", "04 · MCP",
         "05 · Skills and Marketplace", "06 · Subagents",
         "07 · Automation"],
        3)


def slide_04(prs):
    slide = chrome(prs, "PART 1 · 01 · INSTALL AND BASIC USAGE",
                   "Three ways in: editor, CLI, cloud", 4)
    cards = [
        ("1", "Cursor Editor",
         "Download from cursor.com — macOS / Windows / Linux."),
        ("2", "Cursor CLI (agent)",
         "curl … | bash (Unix) · irm … | iex (Windows) — headless-capable."),
        ("3", "Background / Cloud Agents",
         "cursor.com/agents — run tasks in the cloud with no editor open."),
        ("4", "Multi-repo / workspace",
         "Open several repos; distinguish owned context from read-only."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.05 + 1.28 * (i // 2)
        num_card(slide, x, y, 5.835, 1.0, badge, title, desc, 0.45)

    label(slide, 0.7, 4.51, 5.9, "Start in any project")
    code_block(slide, 0.7, 4.85, 5.9, 1.3, [
        "cd your-project",
        "agent               # interactive CLI, login the 1st time",
    ])
    card(slide, 6.9, 4.51, 5.75, 1.55, "The same agent everywhere", [
        "Editor (chat / Agent mode) · CLI (agent)",
        "Web (cursor.com/agents, async tasks)",
        "Bugbot reviews PRs automatically on GitHub",
    ], BLUE, 0.65)
    quote(slide, "Install is trivial; the value is in how you use it.")


def slide_05(prs):
    slide = chrome(prs, "PART 1 · 01 · INSTALL AND BASIC USAGE",
                   "Two modes: interactive and headless", 5)
    card(slide, 0.7, 2.05, 5.85, 2.5, "Interactive — think with you", [
        "Editor (chat/Agent mode) or CLI `agent`",
        "Plan mode: proposes a plan you approve before anything is touched",
        "Ideal for exploring, designing, debugging",
    ], ACCENT, 1.6)
    card(slide, 6.8, 2.05, 5.85, 2.5, "Headless (agent -p) — scripts and CI", [
        "Print mode: prompt in (arg), result on stdout",
        "Combine with --force / --output-format text|json",
        "The foundation of automation (section 07)",
    ], BLUE, 1.6)
    label(slide, 0.7, 4.71, 11.95, "The same Cursor, two ways to invoke it")
    code_block(slide, 0.7, 5.05, 11.95, 1.1, [
        'agent -p "summarize the changes on this branch"',
        'agent -p --output-format text "security-review files touched vs main"',
    ])


def slide_06(prs):
    slide = chrome(prs, "PART 1 · 02 · MEMORY AND INSTRUCTIONS",
                   "AGENTS.md: the two-level pattern", 6)
    card(slide, 0.7, 2.05, 5.85, 3.0, "Level 1 — always loaded (small)", [
        "Minimum orientation: what the project is, repo map, commands",
        "Only one-line POINTERS to everything else",
        "The agent's 30-second onboarding",
        "Write-once rule: the core points, it does not copy",
    ], ACCENT, 2.1)
    card(slide, 6.8, 2.05, 5.85, 3.0, "Level 2 — on demand (the detail)", [
        "data/changes/STATUS.md — living state per ticket",
        "PLAYBOOK.md — debugging lessons",
        "SHARP_EDGES.md — 'do not touch' invariants",
        "<TICKET>/<TICKET>.md — self-contained ledgers",
    ], BLUE, 2.1)
    quote(slide, "The same cut as in Claude Code: pointers, not copies — "
                 "AGENTS.md stays deliberately small.")


def slide_07(prs):
    slide = chrome(prs, "PART 1 · 02 · INSTRUCTIONS, PERMISSIONS AND MEMORIES",
                   "AGENTS.md, rules, permissions, and memory", 7)
    label(slide, 0.7, 1.81, 6.0, "Rules (.mdc) + AGENTS.md (combine, don't compete)")
    code_block(slide, 0.7, 2.15, 6.0, 2.35, [
        ".cursor/rules/00-core.mdc     # alwaysApply: true",
        ".cursor/rules/01-tools.mdc    # @file = include content",
        "AGENTS.md                     # root (+ nested: most",
        "                               # specific wins), no frontmatter",
        "~/.cursor/permissions.json    # mcpAllowlist server:tool",
    ])
    card(slide, 7.0, 1.82, 5.65, 2.55, "Permissions = allowlist, not UI alone", [
        "permissions.json: mcpAllowlist / terminalAllowlist (glob)",
        "4 rule modes: Always / Apply Intelligently /",
        "  Apply to Specific Files / Apply Manually",
        "The human owns push/PR/deploy",
    ], GREEN, 1.65)
    card(slide, 0.7, 4.6, 11.95, 1.5,
         "Memories — a system DISTINCT from AGENTS.md/rules", [
             "Cursor generates Memories automatically from your chats "
             "(they are not files you write).",
             "Not the same mechanism as Claude Code auto-memory: treat them "
             "as a complement, not a substitute for the two-level pattern.",
         ], BLUE, 0.6)


def slide_08(prs):
    slide = chrome(prs, "PART 1 · 02 · SESSIONS",
                   "Where Cursor lives: editor, CLI, and cloud", 8)
    cards = [
        ("▣", "Editor",
         "Chat + Agent mode + Plan mode; inline diffs, review in the IDE itself."),
        ("›_", "Cursor CLI",
         "agent in the terminal — same as the editor, no UI; ideal for SSH/servers."),
        ("☁", "Background / Cloud Agents",
         "cursor.com/agents runs async tasks; follow progress from the web."),
        ("✓", "Bugbot",
         "Reviews every PR on GitHub/GitLab/Bitbucket automatically, no custom script."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.15 + 1.38 * (i // 2)
        num_card(slide, x, y, 5.835, 1.1, badge, title, desc, 0.55)
    quote(slide, "Same agent, same AGENTS.md/rules/MCP — only where it runs changes.")


def slide_09(prs):
    slide = chrome(prs, "PART 1 · 03 · CONTEXT",
                   "Context window: the resource that governs everything", 9)
    rows = [
        ("System / agent prompt", "Hidden — always first", "fixed"),
        ("Rules alwaysApply + AGENTS.md",
         "You control this — hence the two-level pattern", "you decide"),
        ("Memories", "If any — different system, review what snuck in", "variable"),
        ("Tools MCP", "Schemas of enabled tools", "per server"),
        ("Conversation + results",
         "Files read, command output… grows every turn", "grows"),
    ]
    for i, (name, desc, note) in enumerate(rows):
        stack_row(slide, 2.0 + 0.58 * i, name, desc, note)

    label(slide, 8.6, 2.01, 4.05, "The controls (CLI + editor UI)")
    code_block(slide, 8.6, 2.35, 4.05, 2.55, [
        "agent> /summarize  # free context",
        "agent> /rewind     # previous message",
        ("# new chat / new agent invocation", MUTED),
        ("# + context ring in the editor", MUTED),
    ])
    line(slide, 0.7, 5.05, 11.95, 0.9, [
        R("Performance degrades BEFORE the window fills: ", 12, TEXT, True),
        R("noisy context = worse decisions. Cursor auto-summarizes as it "
          "nears the limit (plus manual /summarize) — do not assume exact "
          "parity with Claude Code's /compact.", 12, MUTED),
    ], ls=1.1)
    quote(slide, "The editor's context ring tells you which block is swelling — "
                 "measure before you cut.")


def slide_10(prs):
    slide = chrome(prs, "PART 1 · 03 · CONTEXT",
                   "Context hygiene: maximum signal per token", 10)
    cards = [
        ("1", "Minimal AGENTS.md / rules",
         "If you can delete it without the agent going wrong, delete it. "
         "Two levels + pointers."),
        ("2", "@file only when needed",
         ".mdc rules support @file to include content — use carefully, "
         "it is eager load."),
        ("3", "MCP in moderation",
         "Each server adds its tool block. Disable what the project does not use."),
        ("4", "Subagents for investigation",
         "Messy exploration goes to a subagent (section 06); only the summary returns."),
        ("5", "Plan → Agent",
         "Separate exploration (Plan mode) from implementation; noise does not stay."),
        ("6", "Targeted reads",
         "“read config/auth.ts” > “understand auth”. Preview: CodeGraph-first "
         "(Part 2)."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.0 + 1.46 * (i // 2)
        num_card(slide, x, y, 5.835, 1.18, badge, title, desc, 0.63)
    quote(slide, "Part 2's tool prevalence is, at bottom, "
                 "context policy.")


def slide_11(prs):
    slide = chrome(prs, "PART 1 · 03 · CONTEXT",
                   "Prompt caching: don't pay the same twice", 11)
    line(slide, 0.7, 1.9, 11.95, 0.55, [
        R("This is from the Anthropic API, not Cursor — but if you automate with the "
          "Cursor SDK against Claude, it applies the same. Strict order ", 12.5, MUTED),
        R("Tools → System → Messages", 12.5, ACCENT, True, MONO),
        R(": a change invalidates its level and the ones after.", 12.5, MUTED),
    ])
    for i, (ttl, write, read) in enumerate([
        ("TTL 5 min (default)", "write 1.25×", "read 0.1×"),
        ("TTL 1 hour", "write 2×", "read 0.1×"),
    ]):
        y = 2.6 + 0.6 * i
        panel(slide, 0.7, y, 5.85, 0.5)
        line(slide, 0.95, y + 0.05, 2.5, 0.4, [R(ttl, 11.5, BLUE, True)],
             anchor=MSO_ANCHOR.MIDDLE)
        line(slide, 3.5, y + 0.05, 1.7, 0.4, [R(write, 11.5, TEXT)],
             anchor=MSO_ANCHOR.MIDDLE)
        line(slide, 5.2, y + 0.05, 1.2, 0.4, [R(read, 11.5, GREEN, True)],
             anchor=MSO_ANCHOR.MIDDLE)

    tb(slide, 0.7, 3.95, 5.85, 1.6, [
        [R("A 50-turn session rereads the prefix 50 times at 0.1× — ", 12, TEXT),
         R("that discount is what makes long sessions viable.", 12, GREEN, True)],
        [R("Minimum cacheable ~1k tokens · max 4 breakpoints · diagnose via "
           "usage.cache_read_input_tokens.", 10.5, MUTED)],
    ], sa=8, ls=1.1)

    label(slide, 6.8, 2.26, 5.85, "Stable first; breakpoint at the end of the stable part")
    code_block(slide, 6.8, 2.6, 5.85, 2.9, [
        'system=[{ "type": "text",',
        '  "text": STABLE_INSTRUCTIONS,   # stable',
        '  "cache_control":',
        '    {"type": "ephemeral"} }]   # breakpoint',
        'messages=[{"role": "user",',
        '  "content": query }]  # variable: OUTSIDE',
        ("# classic error: timestamp BEFORE the", MUTED),
        ("# breakpoint -> 0 hits and nobody knows why", MUTED),
    ])
    quote(slide, "Runnable demo: ejemplos/prompt-caching/cache_demo.py — talks to the "
                 "Anthropic API, not the Cursor product.")


def slide_12(prs):
    slide = chrome(prs, "PART 1 · 03 · CONTEXT",
                   "Caching: what you control in Cursor and what you don't", 12)
    rows = [
        (True, "Small, stable AGENTS.md + rules",
         "Short prefix that does not change between turns → better behavior"),
        (False, "Edit rules/AGENTS.md mid-session",
         "Base context changes → pay the tax again"),
        (False, "Many active MCP servers",
         "Large, changing tool block → more fixed context"),
        (True, "New session / clean chat between unrelated tasks",
         "Avoid dragging an infinite transcript"),
        (False, "Assume Cursor exposes cache_control like the API",
         "No — different product; don't copy Claude Code env vars"),
    ]
    for i, (verdict, headline, detail) in enumerate(rows):
        flag_row(slide, 2.05 + 0.68 * i, verdict, headline, detail)
    line(slide, 0.7, 5.6, 11.95, 0.55, [
        R("Cursor does not expose cache_control as a product: you get what the model "
          "provider decides. What you do control: size/stability of AGENTS.md + rules "
          "+ MCP tools.", 11.5, MUTED)])
    quote(slide, "Lean, stable context performs better in any product — Cursor "
                 "included — even if you never see the discount on screen.")


def slide_13(prs):
    slide = chrome(prs, "PART 1 · 04 · MCP",
                   "Model Context Protocol: connect your world", 13)
    line(slide, 0.7, 1.9, 11.95, 0.6, [
        R("Same open standard as in Claude Code — CodeGraph, Serena, Playwright "
          "speak MCP in both. Project config in ", 12.5, MUTED),
        R(".cursor/mcp.json", 12.5, ACCENT, True, MONO),
        R(".", 12.5, MUTED),
    ])
    cards = [
        ("P", "Project scope", ".cursor/mcp.json — versioned, shared with the team."),
        ("U", "User scope",
         "~/.cursor/mcp.json — edit directly, or via Settings → MCP."),
        ("✓", "Permissions", "permissions.json: mcpAllowlist / terminalAllowlist."),
        ("★", "Day-to-day servers", "serena · context7 · playwright · codegraph (+ optional supabase)."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.785
        y = 2.55 + 1.07 * (i // 2)
        num_card(slide, x, y, 5.865, 0.85, badge, title, desc, 0.3)

    label(slide, 0.7, 4.62, 11.95,
          "Add a server (secrets via env var; reload required)")
    code_block(slide, 0.7, 4.96, 11.95, 1.1, [
        '"serena": { "command": "uvx", "args": ["--from", '
        '"git+https://github.com/oraios/serena", "serena", "start-mcp-server"] }',
        ("# after editing .cursor/mcp.json: reload/restart Cursor — no hot-reload",
         MUTED),
    ])


def slide_14(prs):
    slide = chrome(prs, "PART 1 · 05 · SKILLS AND MARKETPLACE",
                   "Skills: compatible with Claude Code", 14)
    cards = [
        ("1", "Tools",
         "What the agent can do: Read/Edit/Shell/Grep + Task + mcp__*. "
         "Governed by permissions.json (mcpAllowlist)."),
        ("2", "Skills",
         ".cursor/skills/<n>/SKILL.md with frontmatter name+description. "
         "Auto-selected by description."),
        ("3", "Interop with Claude Code",
         "Cursor also reads .claude/skills/ — the SAME SKILL.md works in both "
         "products."),
        ("4", "Marketplace",
         "cursor.com/marketplace: installable packages of skills+subagents+MCP+hooks"
         "+rules."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.05 + 1.43 * (i // 2)
        num_card(slide, x, y, 5.835, 1.15, badge, title, desc, 0.6)

    line(slide, 0.7, 4.75, 11.95, 0.9, [
        R("What is NO LONGER a gap vs Claude Code: ", 12.5, GREEN, True),
        R("when the original adaptation guide was written, Cursor had neither Skills nor "
          "Marketplace. Both are real today — check docs.cursor.com before "
          "assuming something “has no equivalent”.", 12.5, MUTED),
    ], ls=1.05)
    quote(slide, "Skill = capability Cursor chooses by its description. "
                 "Marketplace = the versioned distribution.")


def slide_15(prs):
    slide = chrome(prs, "PART 1 · 06 · SUBAGENTS",
                   "Subagents: noise dies outside your session", 15)
    card(slide, 0.7, 2.0, 5.85, 2.2, "Native Cursor feature", [
        "Task-style delegation: isolated context per subagent",
        "Invoked in natural language (or /name)",
        "Parallel execution for independent work",
        "Base types documented loosely — do not assume exact names",
    ], BLUE, 1.3)
    label(slide, 6.8, 2.01, 5.85, 'Custom = .cursor/agents/<name>.md')
    code_block(slide, 6.8, 2.35, 5.85, 2.2, [
        ("# .cursor/agents/refactor-scout.md", MUTED),
        "---",
        "name: refactor-scout",
        "description: Scout blast radius before rename",
        "readonly: true",
        "---",
        "Use CodeGraph and THEN Serena…",
    ])
    card(slide, 0.7, 4.75, 5.85, 1.5, "What you need to know", [
        "ISOLATED context: only the summary returns to your session",
        "Does NOT inherit your conversation: give context in the launch prompt",
    ], ACCENT, 0.6)
    quote(slide, "Ready example: ejemplos/subagents/.cursor/agents/refactor-scout.md "
                 "(also prompts/ if you prefer to paste at launch).")


def slide_16(prs):
    slide = chrome(prs, "PART 1 · 06 · SUBAGENTS",
                   "Real parallelism: multiple Background/Cloud Agents", 16)
    card(slide, 0.7, 2.0, 5.85, 2.6, "What does NOT exist in Cursor", [
        "Claude Code Agent Teams: lead + teammates + shared inbox",
        "No direct messaging between agents",
        "Don't try a 1:1 port — there is no product equivalent",
    ], ACCENT, 1.7)
    label(slide, 6.8, 1.81, 5.85, "The pragmatic substitute")
    code_block(slide, 6.8, 2.15, 5.85, 1.9, [
        "// cursor.com/agents — several tasks in parallel",
        "// each Background/Cloud Agent = full session,",
        "// on its own branch, no coordination between them",
    ])
    card(slide, 6.8, 4.25, 5.85, 1.9, "How it works in practice", [
        "Partition work by branch/file before launching",
        "A human (or the main agent) integrates results",
        "Cost: each Cloud Agent is a full session",
    ], BLUE, 1.0)
    line(slide, 0.7, 4.75, 5.85, 1.3, [
        R("Ask in natural language: ", 11.5, MUTED),
        R("launch three Background Agents, one per module, each on its own branch; "
          "you review and merge.", 11.5, TEXT),
    ], ls=1.1)
    quote(slide, "No shared inbox: partition files/branches — each agent "
                 "owns its own.")


def slide_17(prs):
    slide = chrome(prs, "PART 1 · 06 · SUBAGENTS",
                   "Subagent or Background/Cloud Agent? The decision", 17)
    line(slide, 4.6, 2.0, 3.7, 0.4, [R("SUBAGENT", 13, BLUE, True)],
         align=PP_ALIGN.CENTER)
    line(slide, 8.6, 2.0, 3.7, 0.4, [R("BACKGROUND / CLOUD AGENT", 13, ACCENT, True)],
         align=PP_ALIGN.CENTER)
    rows = [
        ("Context", "Isolated; returns a summary", "Full session, async, on its branch"),
        ("Communication", "Result only → main session",
         "None between agents (no inbox)"),
        ("Cost", "Low (the expensive work dies outside)", "High (N full sessions)"),
        ("Use it for", "Side-quests: investigate, verify",
         "Long/async work, or real parallelism"),
        ("Config", ".cursor/agents/*.md, prompt or skill",
         "cursor.com/agents (UI / handoff &)"),
    ]
    for i, (field, left, right) in enumerate(rows):
        compare_row(slide, 2.45 + 0.62 * i, field, left, right)
    line(slide, 0.7, 5.68, 11.95, 0.5, [
        R("Bridge to Part 2: ", 11.5, GREEN, True),
        R("GSD (Claude Code) packages roles as plugin subagents; in Cursor those "
          "roles live as .cursor/agents/ + skills — this project uses data/changes/, "
          "not GSD (see Part 2).", 11.5, MUTED),
    ])
    quote(slide, "Subagent so noise dies outside; Background/Cloud Agent for "
                 "long or async work. The cost is not the same.")


def slide_18(prs):
    slide = chrome(prs, "PART 1 · 07 · AUTOMATION",
                   "Hooks: deterministic control", 18)
    card(slide, 0.7, 2.0, 5.85, 3.1, "The hook contract", [
        ".cursor/hooks.json (project) or ~/.cursor/hooks.json (user)",
        "Event payload on STDIN, JSON response on STDOUT",
        "Events: beforeShellExecution, beforeMCPExecution, beforeReadFile,",
        "  afterFileEdit, preToolUse/postToolUse, beforeSubmitPrompt, stop…",
        "permission: prefer deny for gates (ask is often ignored)",
        "exit 2 = deny too · failClosed if the hook crashes",
    ], ACCENT, 2.2)
    label(slide, 6.8, 2.01, 5.85, "Blocking pattern (JS)")
    code_block(slide, 6.8, 2.35, 5.85, 2.4, [
        "// beforeReadFile — block .env",
        'const p = JSON.parse(require("fs")',
        '  .readFileSync(0, "utf8"));',
        'if ((p.path || "").includes(".env")) {',
        "  console.log(JSON.stringify(",
        '    {permission:"deny",',
        '     user_message:"Blocked: .env"}));',
        "  process.exit(0);",
        "}",
    ])
    quote(slide, "With AGENTS.md you ask it to behave; with a hook you guarantee it — "
                 "same as Claude Code, different contract.")


def slide_19(prs):
    slide = chrome(prs, "PART 1 · 07 · AUTOMATION",
                   "From hooks to agents in the cloud", 19)
    cards = [
        ("a", "Hooks",
         "Security (.env), formatting, blocking type-check, veto of push/deploy."),
        ("b", "Headless / CLI",
         "agent -p in scripts and CI (print mode; --force / --output-format)."),
        ("c", "CI/CD",
         "Bugbot (native) for PR review, or Cursor SDK in your own GitHub Action."),
        ("d", "Scheduling",
         "Automations: cron + triggers (Slack, Linear, PR merged, PagerDuty)."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.785
        y = 2.05 + 1.17 * (i // 2)
        num_card(slide, x, y, 5.865, 0.95, badge, title, desc, 0.4)

    label(slide, 0.7, 4.34, 11.95,
          "Cursor SDK (@cursor/sdk) — Agent.prompt (one-shot) · Agent.create+send (stream/multi-turn)")
    code_block(slide, 0.7, 4.68, 11.95, 1.3, [
        'import { Agent } from "@cursor/sdk";',
        "const result = await Agent.prompt(prompt, {",
        "  apiKey: process.env.CURSOR_API_KEY!, local: { cwd } });",
        "// multi-turn: Agent.create(...) → agent.send(...) → run.stream()",
    ])


def slide_20(prs):
    part_divider(
        prs, "PART 2", "The methodology",
        "Tool-agnostic: demonstrated with Cursor, but the discipline travels.",
        ["08 · The 11-stage flow + real example",
         "09 · The method's tools",
         "10 · Transfer: from Claude Code to Cursor",
         "11 · Ops: machine-sync + shared record (S3)"],
        20)


def slide_21(prs):
    slide = chrome(prs, "PART 2 · 08 · THE METHODOLOGY",
                   "The agent is a disciplined collaborator", 21)
    panel(slide, 0.7, 2.0, 11.95, 1.15, border=ACCENT, border_w=THICK, fill=BAND)
    line(slide, 1.0, 2.18, 11.4, 0.85,
         [R("“Autonomy is earned decision-by-decision, not granted in bulk.”", 17, TEXT,
            True)],
         anchor=MSO_ANCHOR.MIDDLE)
    card(slide, 0.7, 3.4, 5.85, 2.05, "The agent owns", [
        "Investigation and diagnosis",
        "Plans and implementation",
        "Tests and documentation",
    ], BLUE, 1.15)
    card(slide, 6.8, 3.4, 5.85, 2.05, "The human owns", [
        "Go/no-go decisions and scope",
        "Every external action: push, PR, deploy",
        "Plan approval",
    ], GREEN, 1.15)
    line(slide, 0.72, 5.62, 11.9, 0.5, [
        R("Tool-agnostic:  ", 11.5, ACCENT, True),
        R("this is not a Cursor flow — it is a way of working with ANY agent. "
          "Section 10 shows transferring it from Claude Code (and it also "
          "reaches GitHub Copilot).", 11, MUTED),
    ], ls=1.03)
    quote(slide, "The model does not earn trust for free: it earns it decision by decision, "
                 "with evidence.")


def slide_22(prs):
    slide = chrome(prs, "PART 2 · 08 · THE METHODOLOGY", "The real 11-stage flow", 22)
    steps = [
        ("1", "Orient: skill kg + history-first + status-first", False),
        ("2", "Inbound triage: is the symptom in the contract?", True),
        ("3", "Regression vs. pre-existing (prove it)", True),
        ("4", "Investigate: cheap deterministic oracle", False),
        ("5", "Plan mode → explicit human agreement", True),
        ("6", "Implement: TDD RED → GREEN, minimal", False),
        ("7", "Verify: outbound gate ×5 (instrument → list → image → look)", True),
        ("8", "Document — each thing once", False),
        ("9", "Sanitize — added lines (skill sanitise-diff)", False),
        ("10", "Handoff: the human does push / PR", False),
        ("11", "Bugbot/bot review + persist lessons", True),
    ]
    for i, (number, text, gate) in enumerate(steps):
        x = 0.7 if i < 6 else 6.85
        y = 2.05 + 0.66 * (i if i < 6 else i - 6)
        flow_step(slide, x, y, number, text, gate)
    line(slide, 0.7, 6.02, 11.95, 0.3, [
        R("■ ", 11, ACCENT, True),
        R("Coral boxes are GATES (decision points). A red gate = STOP: "
          "do not write code.", 11, MUTED),
    ])
    quote(slide, "The agent orchestrates and is where inference lives; the expensive work "
                 "concentrates in plan/code/verify, not in search.")


def slide_23(prs):
    slide = chrome(prs, "PART 2 · 08 · THE METHODOLOGY",
                   "A real example, end to end", 23)
    panel(slide, 0.7, 1.95, 11.95, 0.62, border=ACCENT, fill=BAND)
    line(slide, 0.95, 2.0, 11.5, 0.52, [
        R("QA bug: ", 12.5, ACCENT, True),
        R("“a field shows empty in the UI, but it is written in the PDF.”", 12.5, TEXT),
    ], anchor=MSO_ANCHOR.MIDDLE)
    rows = [
        ("Orient",
         "Skill kg surfaces the danger zone (a SHARP_EDGE that constrains the fix); "
         "git/gh give the baseline."),
        ("Contract",
         "The field is empty in the endpoint JSON (Playwright MCP) → the Lambda is "
         "responsible."),
        ("Provenance",
         "Re-extract on the prior baseline: already empty → pre-existing, not a regression."),
        ("Investigate",
         "Serena + CodeGraph (MCP) locate the detector; a deterministic _diag_pdf.py "
         "gives the cause: NO LLM."),
        ("Fix + verify",
         "RED test → structural fix (not client) → no-op regression + local contract "
         "(wrapper) + inside the deployed image."),
        ("Handoff",
         "Skill sanitise-diff → human does push/PR. Bugbot finds one more case → "
         "test + PLAYBOOK."),
    ]
    for i, (label_text, detail) in enumerate(rows):
        step_row(slide, 2.78 + 0.56 * i, label_text, detail)
    quote(slide, "The model does not earn trust for free: it earns it decision by decision, "
                 "with evidence.")


def slide_24(prs):
    slide = chrome(prs, "PART 2 · 09 · METHOD TOOLS",
                   "Which tool the agent uses, and when", 24)
    line(slide, 0.7, 1.9, 11.95, 0.5, [
        R("Rule from .cursor/rules/01-tool-prevalence.mdc: ", 12, MUTED),
        R("codegraph_explore FIRST (survey in 1 call) · Serena "
          "find_referencing_symbols for the precise check before rename. ",
          12, TEXT, True),
        R("Order: cheap → expensive, deterministic → probabilistic.", 12, MUTED),
    ])
    rows = [
        ("Orient",
         "Skill kg (ticket graph → danger zone) · STATUS.md · git · gh",
         "no inference", False),
        ("Navigate (survey)",
         "CodeGraph codegraph_explore (MCP) — source + paths + blast radius + coverage",
         "1 call", False),
        ("Refactor-check",
         "Serena find_referencing_symbols (MCP) — disambiguates by class "
         "(before rename)", "precise", False),
        ("Diagnose", "Deterministic oracle: parser / validator / _diag_*.py",
         "no inference", False),
        ("Environment (logs, config)",
         "AWS CLI — CloudWatch · lambda get-function · SQS/DLQ", "read-only", False),
        ("Output contract", "Playwright MCP against the endpoint the consumer sees",
         "reproduce", False),
        ("Only at the end",
         "The agent's roll — to VERIFY the fix, not to diagnose",
         "metered", True),
    ]
    for i, (phase, tool, note, highlight) in enumerate(rows):
        tool_row(slide, 2.5 + 0.6 * i, phase, tool, note, highlight)


def slide_25(prs):
    slide = chrome(prs, "PART 2 · 09 · METHOD TOOLS",
                   "Each tool in one sentence — and its phase", 25)
    cards = [
        ("CG", "CodeGraph → investigate",
         "CODE graph (tree-sitter→SQLite, local, via MCP). Survey in 1 call: "
         "source + callers + blast radius + coverage."),
        ("Se", "Serena → pre-refactor",
         "Semantic navigation via LSP (MCP). find_referencing_symbols disambiguates by "
         "class: the precise check before rename."),
        ("G", "GSD → plan/execute/verify",
         "The productized method — Claude Code only today. In Cursor: Plan mode + skill "
         "methodology-plan + data/changes/."),
        ("kg", "Skill kg → orient",
         "Project MEMORY graph (with graphify): tickets + sharp edges. "
         "History-first without LLM (Part 3)."),
        ("Pw", "Playwright (MCP) → triage and outbound",
         "Reproduce the symptom where the consumer sees it: the real endpoint, not an "
         "internal function."),
        ("Or", "Oracles → diagnose",
         "Own parsers/validators/_diag_*.py: cheap, reproducible answer before "
         "spending the agent's roll."),
    ]
    for i, (badge, title, desc) in enumerate(cards):
        x = 0.7 if i % 2 == 0 else 6.815
        y = 2.0 + 1.46 * (i // 2)
        num_card(slide, x, y, 5.835, 1.18, badge, title, desc, 0.63)
    quote(slide, "The classic inversion — the model diagnoses — is what this order "
                 "avoids: the model verifies; oracles diagnose.")


def slide_26(prs):
    slide = chrome(prs, "PART 2 · 09 · METHOD TOOLS",
                   "GSD: only exists in Claude Code (today)", 26)
    card(slide, 0.7, 2.0, 5.85, 3.3, "What GSD is", [
        "Claude Code plugin for phased project management",
        "Cycle: discuss → plan (gate) → execute → verify",
        "Subagents gsd-planner, gsd-plan-checker, gsd-executor,",
        "  gsd-code-reviewer, gsd-verifier",
        "Versioned state in .planning/",
        "No official port to Cursor",
    ], ACCENT, 2.4)
    card(slide, 6.8, 2.0, 5.85, 3.3, "Cursor equivalent", [
        "Plan mode for the discuss → plan gate",
        "Skill methodology-plan fills the template before implementing",
        "gsd-* roles → Task/subagents + prompts with that role",
        "State in data/changes/ (same as Claude Code)",
        "GSD shines on multi-phase greenfield; not used daily here",
    ], BLUE, 2.4)
    quote(slide, "Same method, without the plugin: in Cursor, Plan mode + skills replicate "
                 "GSD's gates.")


def slide_27(prs):
    slide = chrome(prs, "PART 2 · 09 · METHOD TOOLS",
                   "CodeGraph: local code intelligence (via MCP)", 27)
    card(slide, 0.7, 2.0, 5.85, 2.4, "tree-sitter → SQLite index (.codegraph/)", [
        "Symbols: functions, classes, routes, components",
        "Edges: calls, imports, inheritance, references",
        "Deterministic (from the AST), no API keys",
        "Returns: source + paths + blast radius + coverage",
    ], BLUE, 1.5)
    label(slide, 6.8, 1.86, 5.85, "CLI + register in Cursor (not 'claude mcp add')")
    code_block(slide, 6.8, 2.2, 5.85, 1.85, [
        "codegraph init      # creates .codegraph/",
        "codegraph sync      # incremental after edits",
        'codegraph explore "<symbol|question>"',
        ("# .cursor/mcp.json:", MUTED),
        '"codegraph": {"command":"codegraph",',
        '  "args":["serve","--path","<repo>","--mcp"]}',
    ])
    panel(slide, 6.8, 4.25, 5.85, 1.1)
    line(slide, 7.05, 4.38, 5.4, 0.95, [
        R("First for navigation:  ", 11, ACCENT, True),
        R("source + callers + blast radius + coverage in 1 query (treat it as ALREADY "
          "read). Serena find_referencing_symbols = precise check (disambiguates by "
          "class).", 11, TEXT),
    ], ls=1.03)
    tb(slide, 0.7, 4.6, 5.85, 1.4, [
        [R("58% fewer tool calls · 22% faster", 13, GREEN, True)],
        [R("(per their benchmarks: nearly eliminates file reads)", 10.5, MUTED)],
    ])
    quote(slide, "One query instead of grep→open→follow-import→repeat. "
                 "Less context, more signal.")


def slide_28(prs):
    slide = chrome(prs, "PART 2 · 09 · METHOD TOOLS",
                   "Serena: semantic navigation, precise check", 28)
    card(slide, 0.7, 2.0, 5.85, 2.6, "Semantic navigation via LSP (MCP)", [
        "Operates on SYMBOLS, not text: IDE precision",
        "find_symbol (body=true) — one method from a 5k-line file",
        "get_symbols_overview — skeleton of a file",
        "find_referencing_symbols — who references, BY CLASS",
        "search_for_pattern · activate_project (multi-repo)",
    ], BLUE, 1.7)
    label(slide, 6.8, 2.01, 5.85, "Install (stdio) — do not use 'claude mcp add'")
    code_block(slide, 6.8, 2.35, 5.85, 1.35, [
        "// .cursor/mcp.json",
        '"serena": {"command":"uvx","args":["--from",',
        '  "git+https://github.com/oraios/serena",',
        '  "serena","start-mcp-server"]}',
    ])
    panel(slide, 6.8, 4.0, 5.85, 1.55, border=ACCENT)
    line(slide, 7.05, 4.15, 5.4, 1.3, [
        R("Why it is MANDATORY pre-rename:  ", 11, ACCENT, True),
        R("CodeGraph's flat impact mixes homonymous methods (Invoice.process vs "
          "Refund.process); Serena disambiguates by class. Complementary, "
          "not rivals.", 11, TEXT),
    ], ls=1.05)
    line(slide, 0.7, 4.85, 5.85, 0.9, [
        R("The refactor-scout template ", 11.5, GREEN, True),
        R("(ejemplos/subagents/prompts/) packages the CodeGraph → Serena → grep "
          "order as a procedure.", 11.5, MUTED),
    ], ls=1.05)
    quote(slide, "CodeGraph answers 'what breaks?'; Serena answers "
                 "'exactly who calls THIS process()?'")


def slide_29(prs):
    slide = chrome(prs, "PART 2 · 10 · TRANSFER",
                   "The real proof: from Claude Code to Cursor", 29)
    card(slide, 0.7, 2.0, 5.85, 2.75, "Travels UNCHANGED (the 5 rules)", [
        "Plan → agree → implement",
        "Verify on the CONSUMER's contract",
        "Solve the general class, not one input",
        "Durable trail: why, what, how it was verified",
        "The human owns the external (merge, deploy)",
    ], GREEN, 1.85)
    card(slide, 6.8, 2.0, 5.85, 2.75, "What DOES change: the surface", [
        "CLAUDE.md (+ hierarchy) → AGENTS.md + .cursor/rules/*.mdc",
        "Skills ~/.claude/skills/ → .cursor/skills/ (same SKILL.md)",
        "Hooks + settings.local.json → .cursor/hooks.json",
        "Plan mode → Plan mode (same discipline)",
        "Subagents/Agent Teams → .cursor/agents/*.md (no Teams)",
        "claude -p (headless) → agent -p (Cursor CLI print mode)",
    ], BLUE, 1.85)
    card(slide, 0.7, 4.95, 11.95, 1.25,
         "The starter kit: CURSOR_ADAPTATION.md (docs/ai-agents-code-methodology/)", [
             "Full mapping + cursor/ (rules, skills, MCP, hooks ready to copy) "
             "+ bootstrap-cursor-repo.ps1",
             "The same pack includes COPILOT_ADAPTATION.md: not a special case — the "
             "discipline reaches a third agent",
         ], ACCENT, 0.35)
    quote(slide, "This same repo is the proof: methodology born in Claude Code, "
                 "running in Cursor with CURSOR_ADAPTATION.md.")


def slide_30(prs):
    slide = chrome(prs, "PART 2 · 11 · OPS",
                   "From transporting (tarball) to sharing (S3)", 30)
    card(slide, 0.7, 2.0, 5.85, 2.55, "A · Bring-up: tarball + USB", [
        "Outbound = full copy; inbound = data/ delta only",
        "USB: manual mount in WSL + verify byte by byte",
        "On the destination: bootstrap + target-setup.sh rebuild the tooling",
        "Still the path to stand up a machine from scratch",
    ], BLUE, 1.65)
    card(slide, 6.8, 2.0, 5.85, 2.55, "B · Day to day: S3-backed record", [
        "Narrow scope: changes/**/*.md + graph (no client data)",
        "Write via sync; read via READ-ONLY mount",
        "Dry-run by default; --delete opt-in (don't delete a teammate's work)",
        "Docs = source of truth; graph derived (one publisher)",
    ], GREEN, 1.65)
    card(slide, 0.7, 4.7, 11.95, 1.5,
         "Agent-specific: machine-local IDENTITY.md", [
             "Each machine declares MACHINE_NAME + MACHINE_ROLE (publisher | contributor)",
             "AGENTS.md points to IDENTITY.md — the session reads its role BEFORE acting "
             "(otherwise a contributor republishes the graph)",
         ], ACCENT, 0.55)
    quote(slide, "When the durable trail moves from one machine to a team, the agent "
                 "must know which machine it is on before acting.")


def slide_31(prs):
    part_divider(
        prs, "PART 3", "The ticket graph",
        "A full case built with graphify — project memory, navigable "
        "from Cursor.",
        ["12a · The problem and the design (spike)",
         "12b · Pipeline and commands (skill kg, kg-refresh)",
         "12c · The real graph, visualized",
         "12d · Usage and how it hooks into the methodology"],
        31)


def slide_32(prs):
    slide = chrome(prs, "PART 3 · 12A · TICKET GRAPH",
                   "The problem, and why graphify", 32)
    card(slide, 0.7, 2.0, 5.85, 2.5, "The problem", [
        "~540 writeup files, sharp edges, runbooks, memory",
        "A 'new' bug almost always has prior context that constrains the fix",
        "Finding it by hand = remembering it exists + grep",
        "The graph makes it EXPLICIT and queryable in 1 call",
    ], ACCENT, 1.6)
    card(slide, 6.8, 2.0, 5.85, 2.5, "The corpus: manifest, not glob", [
        "DIFFABLE manifest.txt: ~116 files, ~196k words",
        "Writeups sst-* (+ deterministic fallback) · hubs · runbooks · memory",
        "Hard exclusions: binaries, repeated handovers, stale copies (payload/)",
        "Density without new knowledge = noise",
    ], BLUE, 1.6)
    panel(slide, 0.7, 4.75, 11.95, 1.35, border=GREEN, fill=BAND)
    line(slide, 1.0, 4.92, 11.4, 1.05, [
        R("The technology is graphify — CodeGraph is only the analogy.  ", 12.5, GREEN,
          True),
        R("Same role (queryable graph before touching anything), different domain (tickets, not "
          "code), different tool. The build pipeline was originally generated in "
          "Claude Code; today query and refresh live as skills kg / kg-refresh "
          "— the same SKILL.md works in Cursor and Claude Code.", 11.5, TEXT),
    ], ls=1.12)
    quote(slide, "Full material (design, scripts, tests, real output): "
                 "docs/knowledge-graph/.")


def slide_33(prs):
    slide = chrome(prs, "PART 3 · 12B · TICKET GRAPH",
                   "Pipeline and skills — kg and kg-refresh", 33)
    label(slide, 0.7, 1.91, 6.35, "Build (skill kg-refresh)")
    code_block(slide, 0.7, 2.25, 6.35, 2.15, [
        "kg_refresh.sh prepare   # manifest -> stage",
        "                        # -> scratch OUTSIDE the repo",
        ("# semantic extraction with the agent:", MUTED),
        "                        # nodes+edges+clustering",
        "kg_refresh.sh finalize  # output/ + leak-check",
    ])
    card(slide, 7.35, 1.92, 5.3, 2.5, "The pieces", [
        "Skills kg and kg-refresh (.cursor/skills/ or .claude/skills/ — same file)",
        "kg_query.sh — wraps graphify explain/path + find",
        "kg_refresh.sh — prepare · finalize · bootstrap · snapshot/restore-memory",
        "build_manifest.py + stage_corpus.py — the corpus",
        "test_kg_*.py — the bookends, TESTED",
    ], BLUE, 1.6)
    panel(slide, 0.7, 4.62, 11.95, 1.5)
    tb(slide, 1.0, 4.78, 11.4, 1.2, [
        [R("The gotcha that holds it up:  ", 12, ACCENT, True),
         R("graphify respects .gitignore and all of data/ is gitignored → running the detector "
           "in place finds 0 files. So the corpus is staged in a scratch outside "
           "the repo and artifacts are copied back.", 11.5, TEXT)],
        [R("That is why kg-refresh is a skill, not only a script:  ", 12, GREEN, True),
         R("the semantic step (extraction, parallel subagents) needs an agent; "
           "the bookends are deterministic.", 11.5, TEXT)],
    ], sa=8, ls=1.12)
    quote(slide, "Staging with provenance: sst-5468__sst-5468.md, hub__STATUS.md, "
                 "memory__x.md — every node traces to its source.")


def slide_34(prs):
    slide = chrome(prs, "PART 3 · 12C · TICKET GRAPH",
                   "The real graph: 507 nodes, 35 communities", 34)
    rect(slide, 0.68, 1.93, 8.84, 4.99, None, BORDER, THICK)
    if not KG_IMAGE.exists():
        raise SystemExit(
            f"Missing {KG_IMAGE.name}: generate it with "
            "`python presentacion/capture_kg_graph.py` (requires playwright)."
        )
    slide.shapes.add_picture(str(KG_IMAGE), In(0.7), In(1.95), In(8.8), In(4.95))

    card(slide, 9.7, 1.95, 2.95, 3.4, "The real output", [
        "507 nodes · 672 edges",
        "35 communities",
        "92% EXTRACTED (reliable)",
        "7% INFERRED (conf. 0.7)",
        "116 files, ~196k words",
    ], GREEN, 2.5)
    tb(slide, 9.7, 5.5, 2.95, 0.8, [
        [R("Interactive graph.html:", 11, BLUE, True)],
        [R("node search, filter by community (vis-network).", 10, MUTED)],
    ], ls=1.05)
    quote(slide, "Communities map to real danger zones; god-nodes are "
                 "the free onboarding list.")


def slide_35(prs):
    slide = chrome(prs, "PART 3 · 12D · TICKET GRAPH",
                   "How it is used — and where it hooks in", 35)
    label(slide, 0.7, 1.91, 6.35, "Query: ZERO LLM (kg_query.sh reads graph.json)")
    code_block(slide, 0.7, 2.25, 6.35, 2.0, [
        "kg_query.sh explain <ticket|topic>  # neighbors",
        "kg_query.sh path <A> <B>            # path",
        "kg_query.sh find <substr>           # exact name",
        ("# skill kg-refresh                 # rebuild (cheap)", MUTED),
    ])
    card(slide, 7.35, 1.92, 5.3, 2.5, "Where it hooks in", [
        "Stage 1 (Orient) of the 11-stage flow",
        "history-first rule in .cursor/rules/00-methodology-core.mdc:",
        "skill kg <ticket|topic> BEFORE grep",
        "Points to WHAT to read; does not replace reading",
    ], ACCENT, 1.6)
    line(slide, 0.7, 4.45, 6.35, 0.9, [
        R("Real example: kg get_letter_end ", 12, GREEN, True),
        R("→ the full danger zone instantly: the 5–6 tickets that share that "
          "code.", 11.5, MUTED),
    ], ls=1.1)
    card(slide, 0.7, 5.2, 11.95, 1.0, "Honesty and lifecycle", [
        "Recall in dense zones · EXTRACTED = reliable, INFERRED = hint · internal "
        "(data/) · derived: never travels, rebuild it",
    ], BLUE, 0.1)
    quote(slide, "One semantic step at build time, zero LLM at query time. The graph is "
                 "the map; the agent, the guide.")


def slide_36(prs):
    slide = chrome(prs, "CLOSE", "From assistant to system, in any editor", 36)
    line(slide, 0.7, 1.95, 5.85, 0.4, [R("PART 1 · the tool", 13, ACCENT, True)])
    left = [
        ("1", "Install + memory",
         "Editor, CLI (agent), or cloud; two-level AGENTS.md + rules; Memories separate."),
        ("2", "Context & caching",
         "The budget and the discount: lean and stable wins in both, even if you don't see "
         "it on screen."),
        ("3", "MCP + skills + Marketplace",
         "Connect your world; package and distribute flows — and share SKILL.md with "
         "Claude Code."),
        ("4", "Subagents + automation",
         "Scale the work; hooks that guarantee quality; Bugbot and Automations in the "
         "cloud."),
    ]
    for i, (number, title, desc) in enumerate(left):
        summary_card(slide, 0.7, 2.35 + 0.95 * i, number, title, desc, ACCENT)

    line(slide, 6.8, 1.95, 5.85, 0.4,
         [R("PART 2 + 3 · the method and the graph", 13, BLUE, True)])
    right = [
        ("5", "The 11-stage flow",
         "Deterministic gates; the human owns the decisions."),
        ("6", "The method's tools",
         "CodeGraph · Serena · oracles: cheap→expensive; GSD = the productized method "
         "(Claude Code)."),
        ("7", "Transferable, even in ops",
         "CURSOR_ADAPTATION.md; machine-sync + S3 record with per-machine identity."),
        ("8", "The ticket graph (graphify)",
         "507 nodes · 35 communities: history-first without LLM, skill kg in any "
         "agent."),
    ]
    for i, (number, title, desc) in enumerate(right):
        summary_card(slide, 6.8, 2.35 + 0.95 * i, number, title, desc, BLUE)

    line(slide, 0.7, 6.35, 12.0, 0.4, [
        R("References:  cursor.com/docs  ·  docs.cursor.com/hooks  ·  "
          "cursor.com/marketplace  ·  github.com/colbymchenry/codegraph  ·  "
          "github.com/oraios/serena", 11, DIM)])


SLIDES = [
    slide_01, slide_02, slide_03, slide_04, slide_05, slide_06, slide_07, slide_08,
    slide_09, slide_10, slide_11, slide_12, slide_13, slide_14, slide_15, slide_16,
    slide_17, slide_18, slide_19, slide_20, slide_21, slide_22, slide_23, slide_24,
    slide_25, slide_26, slide_27, slide_28, slide_29, slide_30, slide_31, slide_32,
    slide_33, slide_34, slide_35, slide_36,
]


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    for make in SLIDES:
        make(prs)
    prs.save(str(OUT))
    return len(SLIDES)


if __name__ == "__main__":
    n = build()
    print(f"OK  {OUT}  ({n} slides)")
